from __future__ import annotations

import json
import shutil
import tempfile
import uuid
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "reports" / "TCAD_ML_출품보고서_초안.docx"
FIGURE_DIR = ROOT / "reports" / "source" / "figures"
TRAINING = ROOT / "data" / "training" / "tcad_training_master.xlsx"
SHORT_METRICS = ROOT / "data" / "inverse_test" / "metrics" / "evaluation_short.json"
LONG_METRICS = ROOT / "data" / "inverse_test" / "metrics" / "evaluation_long.json"

# narrative_proposal preset + named override: Korean competition report on A4.
# Full Korean-capable Pretendard static faces are embedded for Word portability.
FONT = "Pretendard"
BLUE = "2457A6"
NAVY = "17365D"
INK = "1F2937"
MUTED = "667085"
LIGHT_BLUE = "EAF2FB"
LIGHT_GRAY = "F4F6F9"
PALE_GOLD = "FFF5D6"
RED = "A61B1B"
WHITE = "FFFFFF"
TABLE_WIDTH_DXA = 9638  # A4 width 210 mm - 20 mm margins on each side.
TABLE_INDENT_DXA = 120


def embed_font(
    docx_path: Path,
    regular_path: Path,
    font_name: str,
    *,
    bold_path: Path | None = None,
    alt_name: str | None = None,
) -> None:
    """Embed Korean-capable regular and bold TrueType faces in the DOCX."""

    with tempfile.TemporaryDirectory() as temp_dir:
        temp_root = Path(temp_dir)
        with zipfile.ZipFile(docx_path, "r") as source:
            source.extractall(temp_root)

        fonts_dir = temp_root / "word" / "fonts"
        fonts_dir.mkdir(parents=True, exist_ok=True)

        rels_path = temp_root / "word" / "_rels" / "fontTable.xml.rels"
        rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
        if rels_path.exists():
            rel_tree = ET.parse(rels_path)
            rel_root = rel_tree.getroot()
        else:
            rel_root = ET.Element(f"{{{rel_ns}}}Relationships")
            rel_tree = ET.ElementTree(rel_root)
        existing_ids = {node.attrib.get("Id", "") for node in rel_root}
        embedded_faces = [
            ("embedRegular", regular_path, "embeddedRegular.odttf"),
        ]
        if bold_path is not None:
            embedded_faces.append(("embedBold", bold_path, "embeddedBold.odttf"))

        embed_elements: list[str] = []
        for embed_element, font_path, embedded_name in embedded_faces:
            key_uuid = uuid.uuid4()
            font_key = "{" + str(key_uuid).upper() + "}"
            key = key_uuid.bytes[::-1]
            font_data = bytearray(font_path.read_bytes())
            for index in range(min(32, len(font_data))):
                font_data[index] ^= key[index % 16]
            (fonts_dir / embedded_name).write_bytes(font_data)

            rel_id_number = 1
            while f"rId{rel_id_number}" in existing_ids:
                rel_id_number += 1
            rel_id = f"rId{rel_id_number}"
            existing_ids.add(rel_id)
            ET.SubElement(
                rel_root,
                f"{{{rel_ns}}}Relationship",
                {
                    "Id": rel_id,
                    "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font",
                    "Target": f"fonts/{embedded_name}",
                },
            )
            embed_elements.append(
                f'<w:{embed_element} r:id="{rel_id}" w:fontKey="{font_key}" '
                'w:subsetted="false"/>'
            )
        rels_path.parent.mkdir(parents=True, exist_ok=True)
        rel_tree.write(rels_path, encoding="UTF-8", xml_declaration=True)

        # Preserve python-docx's namespace declarations verbatim. Re-serializing
        # fontTable.xml may drop w14/w15 declarations referenced by mc:Ignorable.
        font_table_path = temp_root / "word" / "fontTable.xml"
        font_xml = font_table_path.read_text(encoding="utf-8")
        alt_name_xml = f'<w:altName w:val="{alt_name}"/>' if alt_name else ""
        font_entry = (
            f'<w:font w:name="{font_name}">'
            + alt_name_xml
            + '<w:charset w:val="81"/>'
            '<w:family w:val="swiss"/><w:pitch w:val="variable"/>'
            + "".join(embed_elements)
            + '</w:font>'
        )
        font_table_path.write_text(
            font_xml.replace("</w:fonts>", font_entry + "</w:fonts>"),
            encoding="utf-8",
        )

        content_types_path = temp_root / "[Content_Types].xml"
        content_xml = content_types_path.read_text(encoding="utf-8")
        if 'Extension="odttf"' not in content_xml:
            default_entry = (
                '<Default Extension="odttf" '
                'ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont"/>'
            )
            content_xml = content_xml.replace("</Types>", default_entry + "</Types>")
            content_types_path.write_text(content_xml, encoding="utf-8")

        temp_docx = docx_path.with_suffix(".embedded.tmp.docx")
        with zipfile.ZipFile(temp_docx, "w", compression=zipfile.ZIP_DEFLATED) as target:
            for file_path in temp_root.rglob("*"):
                if file_path.is_file():
                    target.write(file_path, file_path.relative_to(temp_root).as_posix())
        shutil.move(temp_docx, docx_path)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != TABLE_WIDTH_DXA:
        raise ValueError(f"table widths must total {TABLE_WIDTH_DXA}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, size=None, bold=None, color=INK, italic=None) -> None:
    run.font.name = FONT
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:eastAsia"), FONT)
    fonts.set(qn("w:cs"), FONT)
    fonts.set(qn("w:hint"), "eastAsia")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_font(paragraph, size=10.5, color=INK, bold=None) -> None:
    for run in paragraph.runs:
        set_font(run, size=size, color=color, bold=bold)


def add_body(doc, text: str, *, bold_lead: str | None = None, after=7, keep=False):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.30
    p.paragraph_format.keep_together = keep
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, size=10.5, bold=True, color=NAVY)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2, size=10.5)
    else:
        r = p.add_run(text)
        set_font(r, size=10.5)
    return p


def add_heading(doc, text: str, level: int = 1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, size={1: 16, 2: 13, 3: 11.5}[level], bold=True, color=BLUE if level < 3 else NAVY)
    return p


def set_row_cant_split(row) -> None:
    """Keep a one-row callout box together on the same page."""
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_callout(doc, title: str, text: str, fill=LIGHT_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_row_cant_split(table.rows[0])
    set_table_geometry(table, [TABLE_WIDTH_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_font(r, size=10.5, bold=True, color=accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.22
    r2 = p2.add_run(text)
    set_font(r2, size=9.8, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int], font_size=9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(value))
        set_font(r, size=font_size, bold=True, color=WHITE)
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            if len(table.rows) % 2 == 1:
                set_cell_shading(cells[i], LIGHT_GRAY)
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            set_font(r, size=font_size, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_figure_caption(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_font(r, size=9, color=MUTED, italic=True)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("페이지 ")
    set_font(run, size=8.5, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def create_workflow_figure(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    font_path = ROOT / "reports" / "assets" / "Pretendard-M.ttf"
    image = Image.new("RGB", (2376, 576), "white")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(str(font_path), 28)
    sub_font = ImageFont.truetype(str(font_path), 21)
    labels = [
        ("TCAD 데이터 생성", "공정변수 X → 소자특성 Y"),
        ("데이터 검증", "Valid/Invalid 및 Vd 결합"),
        ("RF Surrogate", "분류기 + 4개 회귀기"),
        ("후보 탐색", "Lg별 50만 개 공정후보"),
        ("Top 1 추천", "제약 만족 + 최소 오차"),
        ("TCAD 재검증", "최종 레시피 물리 검증"),
    ]
    xs = [45, 437, 829, 1221, 1613, 2005]
    colors = ["#EAF2FB", "#EEF7F2", "#EAF2FB", "#FFF5D6", "#EAF2FB", "#EEF7F2"]
    for i, ((title, sub), x, fill) in enumerate(zip(labels, xs, colors)):
        box = (x, 155, x + 315, 410)
        draw.rounded_rectangle(box, radius=18, fill=fill, outline="#2457A6", width=3)
        title_box = draw.textbbox((0, 0), title, font=title_font)
        title_w = title_box[2] - title_box[0]
        draw.text((x + (315 - title_w) / 2, 225), title, font=title_font, fill="#17365D")
        sub_box = draw.multiline_textbbox((0, 0), sub, font=sub_font, spacing=4, align="center")
        sub_w = sub_box[2] - sub_box[0]
        draw.multiline_text((x + (315 - sub_w) / 2, 292), sub, font=sub_font, fill="#475467", spacing=4, align="center")
        if i < len(xs) - 1:
            start_x = x + 328
            end_x = xs[i + 1] - 12
            draw.line((start_x, 282, end_x, 282), fill="#667085", width=4)
            draw.polygon([(end_x, 282), (end_x - 15, 272), (end_x - 15, 292)], fill="#667085")
    image.save(path)


def create_random_forest_figure(path: Path) -> None:
    """Create a project-specific Random Forest concept diagram.

    The split values and leaf predictions are illustrative, not exported from a
    fitted estimator. This is stated in the figure caption in the report.
    """

    path.parent.mkdir(parents=True, exist_ok=True)
    font_path = ROOT / "reports" / "assets" / "Pretendard-M.ttf"
    image = Image.new("RGB", (2376, 1120), "white")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(str(font_path), 36)
    node_font = ImageFont.truetype(str(font_path), 28)
    small_font = ImageFont.truetype(str(font_path), 23)
    tiny_font = ImageFont.truetype(str(font_path), 20)

    def centered_text(box, text, font, color="#17365D", spacing=4):
        left, top, right, bottom = box
        bounds = draw.multiline_textbbox((0, 0), text, font=font, spacing=spacing, align="center")
        width = bounds[2] - bounds[0]
        height = bounds[3] - bounds[1]
        draw.multiline_text(
            ((left + right - width) / 2, (top + bottom - height) / 2 - bounds[1]),
            text,
            font=font,
            fill=color,
            spacing=spacing,
            align="center",
        )

    def rounded_box(box, text, *, fill="#EAF2FB", outline="#2457A6", font=node_font):
        draw.rounded_rectangle(box, radius=18, fill=fill, outline=outline, width=4)
        centered_text(box, text, font)

    def arrow(start, end, label=None, label_xy=None):
        draw.line((*start, *end), fill="#667085", width=4)
        ex, ey = end
        draw.polygon([(ex, ey), (ex - 13, ey - 9), (ex - 13, ey + 9)], fill="#667085")
        if label and label_xy:
            draw.text(label_xy, label, font=tiny_font, fill="#475467")

    draw.text((70, 35), "한 개의 SS 회귀 트리가 질문하는 방식", font=title_font, fill="#17365D")
    draw.line((1585, 40, 1585, 1060), fill="#D0D5DD", width=3)
    draw.text((1640, 35), "300개 트리의 결합", font=title_font, fill="#17365D")

    root = (520, 120, 1120, 250)
    left_node = (165, 390, 735, 520)
    right_node = (915, 390, 1485, 520)
    rounded_box(root, "Lg ≥ 0.0775 µm인가?\n(65 nm와 90 nm를 나누는 예)")
    rounded_box(left_node, "Halo dose ≥ 1.0×10¹³ cm⁻²인가?")
    rounded_box(right_node, "Anneal time ≤ 0.30 s인가?")

    # Connect the root to its child questions.
    draw.line((650, 250, 450, 390), fill="#667085", width=4)
    draw.line((990, 250, 1200, 390), fill="#667085", width=4)
    draw.text((505, 305), "아니오", font=tiny_font, fill="#475467")
    draw.text((1090, 305), "예", font=tiny_font, fill="#475467")

    leaves = [
        ((55, 735, 390, 865), "Leaf A\n예측 SS = 128"),
        ((455, 735, 790, 865), "Leaf B\n예측 SS = 92"),
        ((860, 735, 1195, 865), "Leaf C\n예측 SS = 105"),
        ((1260, 735, 1595, 865), "Leaf D\n예측 SS = 170"),
    ]
    for box, text_value in leaves:
        rounded_box(box, text_value, fill="#EEF7F2", outline="#3A7D5D", font=node_font)
    for x1, x2, label, label_x in [
        (300, 222, "아니오", 205),
        (590, 622, "예", 605),
        (1050, 1027, "예", 1020),
        (1350, 1427, "아니오", 1360),
    ]:
        draw.line((x1, 520, x2, 735), fill="#667085", width=4)
        draw.text((label_x, 625), label, font=tiny_font, fill="#475467")

    rounded_box((1680, 140, 2290, 245), "동일한 공정 레시피 X", fill="#FFF5D6", outline="#9A6B00")
    tree_boxes = [
        ((1680, 340, 1970, 455), "Tree 1\n82"),
        ((2000, 340, 2290, 455), "Tree 2\n88"),
        ((1680, 520, 1970, 635), "Tree 3\n86"),
        ((2000, 520, 2290, 635), "... Tree 300\n84"),
    ]
    for box, text_value in tree_boxes:
        rounded_box(box, text_value, fill="#EAF2FB", font=node_font)
    draw.line((1985, 245, 1825, 340), fill="#667085", width=4)
    draw.line((1985, 245, 2145, 340), fill="#667085", width=4)
    draw.line((1985, 245, 1825, 520), fill="#667085", width=4)
    draw.line((1985, 245, 2145, 520), fill="#667085", width=4)
    rounded_box((1750, 790, 2220, 930), "최종 예측\n300개 값의 평균", fill="#EEF7F2", outline="#3A7D5D")
    draw.line((1825, 455, 1880, 790), fill="#667085", width=3)
    draw.line((2145, 455, 2090, 790), fill="#667085", width=3)
    draw.line((1825, 635, 1880, 790), fill="#667085", width=3)
    draw.line((2145, 635, 2090, 790), fill="#667085", width=3)
    centered_text((1640, 950, 2330, 1050), "분류기는 평균 확률, 회귀기는 평균 예측값을 사용", small_font, color="#475467")
    draw.text((70, 1015), "※ 질문 순서·기준값·Leaf 값은 학습 데이터의 오차를 줄이는 방향으로 자동 결정된다.", font=small_font, fill="#667085")
    image.save(path)


def create_two_stage_model_figure(path: Path) -> None:
    """Create a training/inference diagram for the classifier-regressor bundle."""

    path.parent.mkdir(parents=True, exist_ok=True)
    font_path = ROOT / "reports" / "assets" / "Pretendard-M.ttf"
    image = Image.new("RGB", (2376, 1040), "white")
    draw = ImageDraw.Draw(image)
    lane_font = ImageFont.truetype(str(font_path), 34)
    box_font = ImageFont.truetype(str(font_path), 25)
    small_font = ImageFont.truetype(str(font_path), 21)
    note_font = ImageFont.truetype(str(font_path), 22)

    def centered_text(box, text, font=box_font, color="#17365D", spacing=4):
        left, top, right, bottom = box
        bounds = draw.multiline_textbbox((0, 0), text, font=font, spacing=spacing, align="center")
        width = bounds[2] - bounds[0]
        height = bounds[3] - bounds[1]
        draw.multiline_text(
            ((left + right - width) / 2, (top + bottom - height) / 2 - bounds[1]),
            text,
            font=font,
            fill=color,
            spacing=spacing,
            align="center",
        )

    def box(coords, text, fill="#EAF2FB", outline="#2457A6", font=box_font):
        draw.rounded_rectangle(coords, radius=18, fill=fill, outline=outline, width=4)
        centered_text(coords, text, font=font)

    def arrow(x1, y1, x2, y2, color="#667085"):
        draw.line((x1, y1, x2, y2), fill=color, width=4)
        if abs(x2 - x1) >= abs(y2 - y1):
            direction = 1 if x2 > x1 else -1
            draw.polygon([(x2, y2), (x2 - 14 * direction, y2 - 9), (x2 - 14 * direction, y2 + 9)], fill=color)
        else:
            direction = 1 if y2 > y1 else -1
            draw.polygon([(x2, y2), (x2 - 9, y2 - 14 * direction), (x2 + 9, y2 - 14 * direction)], fill=color)

    # Training lane.
    draw.text((55, 35), "학습 단계: 실제 Valid 정답으로 두 종류의 모델을 따로 학습", font=lane_font, fill="#17365D")
    box((70, 125, 390, 265), "전체 TCAD\n공정 레시피")
    box((510, 125, 870, 265), "Recipe-level label\nValid / Invalid", fill="#FFF5D6", outline="#9A6B00")
    arrow(390, 195, 510, 195)

    box((1030, 70, 1460, 210), "Valid + Invalid 전체\n분류기 학습", fill="#EAF2FB")
    box((1030, 260, 1460, 400), "실제 Valid이며\n4개 출력이 완전한 자료", fill="#EEF7F2", outline="#3A7D5D")
    draw.line((870, 195, 945, 195), fill="#667085", width=4)
    draw.line((945, 195, 945, 140), fill="#667085", width=4)
    arrow(945, 140, 1030, 140)
    draw.line((945, 195, 945, 330), fill="#667085", width=4)
    arrow(945, 330, 1030, 330)

    box((1630, 70, 2260, 210), "Random Forest Classifier\nP(valid) 학습", fill="#EAF2FB")
    box((1630, 260, 2260, 400), "4개 Random Forest Regressor\nVth · SS · Ion · Ioff", fill="#EEF7F2", outline="#3A7D5D")
    arrow(1460, 140, 1630, 140)
    arrow(1460, 330, 1630, 330)

    # Inference lane.
    draw.line((55, 485, 2320, 485), fill="#D0D5DD", width=3)
    draw.text((55, 525), "추천 단계: classifier를 먼저 통과한 후보만 회귀 예측", font=lane_font, fill="#17365D")
    box((70, 630, 390, 770), "후보 공정조건 X\n(50만 개)")
    box((520, 630, 920, 770), "Classifier\nP(valid) 계산", fill="#FFF5D6", outline="#9A6B00")
    arrow(390, 700, 520, 700)

    box((1080, 550, 1470, 670), "P(valid) < 0.5\n후보 제외", fill="#FDECEC", outline="#A61B1B")
    box((1080, 735, 1470, 855), "P(valid) ≥ 0.5\n회귀기로 전달", fill="#EEF7F2", outline="#3A7D5D")
    draw.line((920, 700, 1000, 700), fill="#667085", width=4)
    draw.line((1000, 700, 1000, 610), fill="#667085", width=4)
    arrow(1000, 610, 1080, 610)
    draw.line((1000, 700, 1000, 795), fill="#667085", width=4)
    arrow(1000, 795, 1080, 795)

    box((1630, 735, 2010, 855), "4개 Regressor\n소자특성 예측", fill="#EAF2FB")
    box((2090, 735, 2305, 855), "목표와 비교\n점수 계산", fill="#EEF7F2", outline="#3A7D5D", font=small_font)
    arrow(1470, 795, 1630, 795)
    arrow(2010, 795, 2090, 795)

    box((125, 900, 2250, 1000), "성능이 나쁜 소자 ≠ Invalid  |  TCAD와 extraction이 정상이라면 큰 SS·낮은 Ion도 회귀 학습에 포함", fill="#FFF5D6", outline="#9A6B00", font=note_font)
    image.save(path)


def create_validity_preprocessing_figure(path: Path, tcad_path: Path, excel_path: Path) -> None:
    """Combine the actual TCAD failure and Excel preprocessing screenshots."""

    path.parent.mkdir(parents=True, exist_ok=True)
    font_path = ROOT / "reports" / "assets" / "Pretendard-M.ttf"
    title_font = ImageFont.truetype(str(font_path), 34)
    note_font = ImageFont.truetype(str(font_path), 25)
    canvas = Image.new("RGB", (2200, 1580), "white")
    draw = ImageDraw.Draw(canvas)

    def paste_scaled(source_path: Path, top: int, max_width: int, max_height: int) -> int:
        source = Image.open(source_path).convert("RGB")
        scale = min(max_width / source.width, max_height / source.height)
        resized = source.resize(
            (int(source.width * scale), int(source.height * scale)),
            Image.Resampling.LANCZOS,
        )
        left = (canvas.width - resized.width) // 2
        canvas.paste(resized, (left, top))
        draw.rectangle(
            (left, top, left + resized.width, top + resized.height),
            outline="#98A2B3",
            width=3,
        )
        return top + resized.height

    draw.text((70, 30), "① Sentaurus Workbench: 빨간색 fail 및 출력 누락 확인", font=title_font, fill="#A61B1B")
    first_bottom = paste_scaled(tcad_path, 90, 2060, 740)

    arrow_y = first_bottom + 28
    draw.line((1100, arrow_y, 1100, arrow_y + 70), fill="#2457A6", width=8)
    draw.polygon(
        [(1100, arrow_y + 90), (1080, arrow_y + 60), (1120, arrow_y + 60)],
        fill="#2457A6",
    )
    note = "동일 공정 레시피와 두 Vd 행을 대조하여 명시적 label로 변환"
    note_box = draw.textbbox((0, 0), note, font=note_font)
    draw.text(((canvas.width - (note_box[2] - note_box[0])) / 2, arrow_y + 17), note, font=note_font, fill="#475467")

    second_title_y = arrow_y + 120
    draw.text((70, second_title_y), "② CSV/Excel 통합·전처리: 정상=Valid 1, 실패·출력 누락=Valid 0", font=title_font, fill="#17365D")
    paste_scaled(excel_path, second_title_y + 60, 2060, 610)
    canvas.save(path)


def create_idvg_doping_comparison(path: Path, abnormal_path: Path, normal_path: Path) -> None:
    """Remove the SVisual control pane and compare curve/profile pairs."""

    path.parent.mkdir(parents=True, exist_ok=True)
    font_path = ROOT / "reports" / "assets" / "Pretendard-M.ttf"
    label_font = ImageFont.truetype(str(font_path), 31)
    panel_font = ImageFont.truetype(str(font_path), 23)
    canvas = Image.new("RGB", (2376, 1540), "white")
    draw = ImageDraw.Draw(canvas)

    def paste_fit(source: Image.Image, box: tuple[int, int, int, int]) -> None:
        left, top, right, bottom = box
        width = right - left
        height = bottom - top
        scale = min(width / source.width, height / source.height)
        resized = source.resize(
            (int(source.width * scale), int(source.height * scale)),
            Image.Resampling.LANCZOS,
        )
        paste_left = left + (width - resized.width) // 2
        paste_top = top + (height - resized.height) // 2
        canvas.paste(resized, (paste_left, paste_top))
        draw.rectangle((left, top, right, bottom), outline="#98A2B3", width=3)

    rows = [
        (
            Image.open(abnormal_path).convert("RGB"),
            "A. 비정상 gate-response 사례",
            "#A61B1B",
            (10, 70, 875, 795),
            (1510, 50, 2320, 795),
        ),
        (
            Image.open(normal_path).convert("RGB"),
            "B. 정상적인 turn-on 사례",
            "#2457A6",
            (5, 20, 880, 750),
            (1510, 10, 2310, 750),
        ),
    ]

    for row_index, (source, label, color, graph_crop, profile_crop) in enumerate(rows):
        row_top = 30 + row_index * 750
        draw.rounded_rectangle(
            (55, row_top, 2321, row_top + 705),
            radius=18,
            fill="#F8FAFC",
            outline="#D0D5DD",
            width=3,
        )
        draw.ellipse((85, row_top + 28, 117, row_top + 60), fill=color)
        draw.text((135, row_top + 24), label, font=label_font, fill=color)
        draw.text((470, row_top + 69), "ID–VG (log scale)", font=panel_font, fill="#475467")
        draw.text((1670, row_top + 69), "NetActive 분포", font=panel_font, fill="#475467")

        graph = source.crop(graph_crop)
        profile = source.crop(profile_crop)
        paste_fit(graph, (90, row_top + 105, 1128, row_top + 670))
        paste_fit(profile, (1248, row_top + 105, 2286, row_top + 670))

    canvas.save(path)


def create_inverse_design_figure(path: Path) -> None:
    """Explain why inverse design is implemented as forward-model search."""

    path.parent.mkdir(parents=True, exist_ok=True)
    font_path = ROOT / "reports" / "assets" / "Pretendard-M.ttf"
    image = Image.new("RGB", (2376, 780), "white")
    draw = ImageDraw.Draw(image)
    lane_font = ImageFont.truetype(str(font_path), 32)
    box_font = ImageFont.truetype(str(font_path), 24)
    note_font = ImageFont.truetype(str(font_path), 21)

    def centered_text(box, text, font=box_font, color="#17365D", spacing=4):
        left, top, right, bottom = box
        bounds = draw.multiline_textbbox((0, 0), text, font=font, spacing=spacing, align="center")
        width = bounds[2] - bounds[0]
        height = bounds[3] - bounds[1]
        draw.multiline_text(
            ((left + right - width) / 2, (top + bottom - height) / 2 - bounds[1]),
            text,
            font=font,
            fill=color,
            spacing=spacing,
            align="center",
        )

    def box(coords, text, fill="#EAF2FB", outline="#2457A6", font=box_font):
        draw.rounded_rectangle(coords, radius=16, fill=fill, outline=outline, width=4)
        centered_text(coords, text, font=font)

    def arrow(x1, y1, x2, y2, color="#667085"):
        draw.line((x1, y1, x2, y2), fill=color, width=4)
        draw.polygon([(x2, y2), (x2 - 15, y2 - 9), (x2 - 15, y2 + 9)], fill=color)

    draw.text((55, 30), "직접 Y→X 회귀를 사용하지 않은 이유", font=lane_font, fill="#A61B1B")
    box((70, 105, 390, 225), "목표 스펙 Y*", fill="#FFF5D6", outline="#9A6B00")
    box((520, 105, 910, 225), "하나의 Y→X\n역회귀 모델", fill="#FDECEC", outline="#A61B1B")
    box((1040, 105, 1390, 225), "평균 레시피 X", fill="#FDECEC", outline="#A61B1B")
    arrow(390, 165, 520, 165)
    arrow(910, 165, 1040, 165)
    centered_text(
        (1490, 85, 2290, 245),
        "동일한 Y를 만드는 X가 여러 개\n→ 정답 하나를 강제로 지정해야 함\n→ 평균 X가 실제 valid recipe라는 보장 없음",
        font=note_font,
        color="#A61B1B",
    )

    draw.line((55, 300, 2320, 300), fill="#D0D5DD", width=3)
    draw.text((55, 335), "현재 방식: 학습된 X→Y forward surrogate를 목적함수처럼 사용", font=lane_font, fill="#17365D")
    steps = [
        ((55, 440, 315, 570), "목표 Y*·허용범위\n이산 Lg 선택", "#FFF5D6", "#9A6B00"),
        ((380, 440, 730, 570), "학습 범위 내부\n후보 X 500,000개", "#EAF2FB", "#2457A6"),
        ((795, 440, 1095, 570), "Classifier\nP(valid)≥0.5", "#EEF7F2", "#3A7D5D"),
        ((1160, 440, 1490, 570), "4개 Regressor\nX→예측 Ŷ", "#EAF2FB", "#2457A6"),
        ((1555, 440, 1865, 570), "목표범위 필터\n오차 score 계산", "#FFF5D6", "#9A6B00"),
        ((1930, 440, 2225, 570), "최소 score\nTop 1 레시피 X*", "#EEF7F2", "#3A7D5D"),
    ]
    for index, (coords, label, fill, outline) in enumerate(steps):
        box(coords, label, fill=fill, outline=outline, font=note_font)
        if index < len(steps) - 1:
            arrow(coords[2], 505, steps[index + 1][0][0], 505)

    box(
        (260, 635, 2115, 735),
        "X* = arg min S(Ŷ, Y*)  |  탐색은 TCAD가 아니라 빠른 ML surrogate에서 수행하고, 최종 후보만 TCAD로 재검증",
        fill="#F4F6F9",
        outline="#667085",
        font=note_font,
    )
    image.save(path)


def create_ui_usage_figure(path: Path, screenshot_path: Path) -> None:
    """Place the actual UI screenshot under a compact four-step reading guide."""

    path.parent.mkdir(parents=True, exist_ok=True)
    font_path = ROOT / "reports" / "assets" / "Pretendard-M.ttf"
    step_font = ImageFont.truetype(str(font_path), 24)
    canvas = Image.new("RGB", (2376, 1480), "white")
    draw = ImageDraw.Draw(canvas)

    steps = [
        ("1", "소자 영역·이산 Lg 선택", "#EAF2FB", "#2457A6"),
        ("2", "평가지표와 Min–Target–Max 입력", "#FFF5D6", "#9A6B00"),
        ("3", "최적 공정조건 탐색 실행", "#EEF7F2", "#3A7D5D"),
        ("4", "Top 1–3 레시피·예상 스펙 확인", "#F4F6F9", "#667085"),
    ]
    left_margin = 65
    gap = 24
    box_width = (canvas.width - 2 * left_margin - 3 * gap) // 4
    for index, (number, label, fill, outline) in enumerate(steps):
        left = left_margin + index * (box_width + gap)
        right = left + box_width
        draw.rounded_rectangle((left, 28, right, 165), radius=18, fill=fill, outline=outline, width=4)
        draw.ellipse((left + 22, 68, left + 76, 122), fill=outline)
        num_bbox = draw.textbbox((0, 0), number, font=step_font)
        draw.text(
            (left + 49 - (num_bbox[2] - num_bbox[0]) / 2, 94 - (num_bbox[3] - num_bbox[1]) / 2 - num_bbox[1]),
            number,
            font=step_font,
            fill="white",
        )
        text_bbox = draw.textbbox((0, 0), label, font=step_font)
        draw.text(
            (left + 96, 94 - (text_bbox[3] - text_bbox[1]) / 2 - text_bbox[1]),
            label,
            font=step_font,
            fill="#17365D",
        )

    screenshot = Image.open(screenshot_path).convert("RGB")
    max_width = canvas.width - 120
    max_height = canvas.height - 230
    scale = min(max_width / screenshot.width, max_height / screenshot.height)
    resized = screenshot.resize(
        (int(screenshot.width * scale), int(screenshot.height * scale)),
        Image.Resampling.LANCZOS,
    )
    left = (canvas.width - resized.width) // 2
    top = 205
    canvas.paste(resized, (left, top))
    draw.rounded_rectangle(
        (left - 3, top - 3, left + resized.width + 3, top + resized.height + 3),
        radius=10,
        outline="#98A2B3",
        width=4,
    )
    canvas.save(path)


def create_evaluation_design_figure(path: Path) -> None:
    """Contrast model-level forward testing with end-to-end inverse testing."""

    path.parent.mkdir(parents=True, exist_ok=True)
    font_path = ROOT / "reports" / "assets" / "Pretendard-M.ttf"
    image = Image.new("RGB", (2376, 920), "white")
    draw = ImageDraw.Draw(image)
    lane_font = ImageFont.truetype(str(font_path), 32)
    box_font = ImageFont.truetype(str(font_path), 23)
    note_font = ImageFont.truetype(str(font_path), 20)

    def centered_text(coords, text, font=box_font, color="#17365D", spacing=4):
        left, top, right, bottom = coords
        bounds = draw.multiline_textbbox((0, 0), text, font=font, spacing=spacing, align="center")
        width = bounds[2] - bounds[0]
        height = bounds[3] - bounds[1]
        draw.multiline_text(
            ((left + right - width) / 2, (top + bottom - height) / 2 - bounds[1]),
            text,
            font=font,
            fill=color,
            spacing=spacing,
            align="center",
        )

    def box(coords, text, fill="#EAF2FB", outline="#2457A6", font=box_font):
        draw.rounded_rectangle(coords, radius=16, fill=fill, outline=outline, width=4)
        centered_text(coords, text, font=font)

    def arrow(x1, y1, x2, y2, color="#667085"):
        draw.line((x1, y1, x2, y2), fill=color, width=4)
        draw.polygon([(x2, y2), (x2 - 15, y2 - 9), (x2 - 15, y2 + 9)], fill=color)

    draw.text((55, 30), "Forward test: X→Y 예측 모델 자체의 정확도", font=lane_font, fill="#17365D")
    forward = [
        ((70, 115, 390, 245), "보지 않은 20%\n공정조건 X_test", "#EAF2FB", "#2457A6"),
        ((530, 115, 850, 245), "학습된 RF\nforward model", "#EEF7F2", "#3A7D5D"),
        ((990, 115, 1310, 245), "예측 스펙\nŶ=f̂(X_test)", "#EAF2FB", "#2457A6"),
        ((1450, 115, 1770, 245), "TCAD 정답\nY_test", "#FFF5D6", "#9A6B00"),
        ((1910, 115, 2260, 245), "예측–정답 비교\nR²·MAE·RMSE", "#F4F6F9", "#667085"),
    ]
    for index, (coords, label, fill, outline) in enumerate(forward):
        box(coords, label, fill=fill, outline=outline)
        if index < len(forward) - 1:
            arrow(coords[2], 180, forward[index + 1][0][0], 180)

    draw.line((55, 320, 2320, 320), fill="#D0D5DD", width=3)
    draw.text((55, 355), "Inverse test: UI 추천 전체 과정이 목표를 달성하는지 검증", font=lane_font, fill="#17365D")
    inverse = [
        ((55, 455, 330, 585), "UI 목표 스펙\nY*·허용범위", "#FFF5D6", "#9A6B00"),
        ((405, 455, 715, 585), "50만 후보 탐색\n분류·회귀·score", "#EAF2FB", "#2457A6"),
        ((790, 455, 1065, 585), "추천 공정조건\nTop 1 X*", "#EEF7F2", "#3A7D5D"),
        ((1140, 455, 1445, 585), "X*를 실제\nTCAD에 재투입", "#FFF5D6", "#9A6B00"),
        ((1520, 455, 1810, 585), "실제 스펙\nY_TCAD(X*)", "#EEF7F2", "#3A7D5D"),
    ]
    for index, (coords, label, fill, outline) in enumerate(inverse):
        box(coords, label, fill=fill, outline=outline)
        if index < len(inverse) - 1:
            arrow(coords[2], 520, inverse[index + 1][0][0], 520)

    box((1880, 390, 2285, 515), "① 예측 Ŷ(X*)와\n실제 Y_TCAD 비교", fill="#EAF2FB", outline="#2457A6", font=note_font)
    box((1880, 550, 2285, 675), "② 실제 Y_TCAD가\n목표범위인지 확인", fill="#FFF5D6", outline="#9A6B00", font=note_font)
    draw.line((1810, 520, 1840, 520), fill="#667085", width=4)
    draw.line((1840, 520, 1840, 452), fill="#667085", width=4)
    arrow(1840, 452, 1880, 452)
    draw.line((1840, 520, 1840, 612), fill="#667085", width=4)
    arrow(1840, 612, 1880, 612)

    box(
        (250, 740, 2125, 855),
        "Forward 성능이 높아야 inverse 추천도 가능하지만, forward test만으로는 최종 추천 레시피의 목표 달성을 증명할 수 없다.",
        fill="#F4F6F9",
        outline="#667085",
        font=note_font,
    )
    image.save(path)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(1.15)
    section.footer_distance = Cm(1.15)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal._element.rPr.rFonts.set(qn("w:cs"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.30

    heading_tokens = {
        1: (16, 16, 8, BLUE),
        2: (13, 12, 6, BLUE),
        3: (11.5, 8, 4, NAVY),
    }
    for level, (size, before, after, color) in heading_tokens.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style._element.rPr.rFonts.set(qn("w:cs"), FONT)
        style._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("TCAD–ML 기반 Planar NMOS 역설계 | 출품 보고서 초안")
    set_font(run, size=8.5, color=MUTED)
    add_page_number(section.footer.paragraphs[0])


def build_report() -> None:
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    flow_figure = FIGURE_DIR / "system_workflow.png"
    forest_figure = FIGURE_DIR / "random_forest_tcad_example.png"
    two_stage_figure = FIGURE_DIR / "classifier_regressor_structure.png"
    validity_figure = FIGURE_DIR / "tcad_to_excel_valid_preprocessing.png"
    device_curve_figure = FIGURE_DIR / "idvg_doping_normal_abnormal.png"
    inverse_design_figure = FIGURE_DIR / "inverse_design_forward_search.png"
    ui_screenshot = FIGURE_DIR / "ui_recommendation_screen.png"
    ui_usage_figure = FIGURE_DIR / "ui_usage_guide.png"
    evaluation_design_figure = FIGURE_DIR / "forward_inverse_evaluation_design.png"
    tcad_fail_figure = FIGURE_DIR / "tcad_fail_example.png"
    excel_valid_figure = FIGURE_DIR / "excel_valid_preprocessing.png"
    abnormal_curve_raw = FIGURE_DIR / "tcad_abnormal_idvg_doping_raw.png"
    normal_curve_raw = FIGURE_DIR / "tcad_normal_idvg_doping_raw.png"
    create_workflow_figure(flow_figure)
    create_random_forest_figure(forest_figure)
    create_two_stage_model_figure(two_stage_figure)
    create_validity_preprocessing_figure(validity_figure, tcad_fail_figure, excel_valid_figure)
    create_idvg_doping_comparison(device_curve_figure, abnormal_curve_raw, normal_curve_raw)
    create_inverse_design_figure(inverse_design_figure)
    create_ui_usage_figure(ui_usage_figure, ui_screenshot)
    create_evaluation_design_figure(evaluation_design_figure)

    short_metrics = json.loads(SHORT_METRICS.read_text(encoding="utf-8"))
    long_metrics = json.loads(LONG_METRICS.read_text(encoding="utf-8"))
    short_sheet = pd.read_excel(TRAINING, sheet_name="Short", header=2)
    long_sheet = pd.read_excel(TRAINING, sheet_name="Long", header=2)

    doc = Document()
    style_document(doc)
    props = doc.core_properties
    props.title = "TCAD–Machine Learning 기반 Planar NMOS 역설계 및 공정 레시피 추천 시스템"
    props.subject = "출품 보고서 초안"
    props.author = "[팀명/작성자 입력]"
    props.keywords = "TCAD, Sentaurus, NMOS, Random Forest, inverse design, surrogate model"

    # Cover: editorial_cover pattern.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(82)
    p.paragraph_format.space_after = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PROJECT REPORT")
    set_font(r, size=11, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("TCAD–Machine Learning 기반\nPlanar NMOS 역설계 및\n공정 레시피 추천 시스템")
    set_font(r, size=27, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(52)
    r = p.add_run("신뢰성 있는 TCAD 데이터와 Random Forest surrogate model을 활용한\n공정 설계 공간 탐색 및 최종 후보 재검증 프레임워크")
    set_font(r, size=12, color=MUTED)
    add_table(doc, ["구분", "내용"], [
        ["출품 분야", "[출품 분야 입력]"],
        ["팀명", "[팀명 입력]"],
        ["작성자", "[작성자 입력]"],
        ["제출일", "[제출일 입력]"],
    ], [2300, 7338], font_size=10)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Sentaurus Process / Sentaurus Device · Random Forest · Inverse Process Design")
    set_font(r, size=9.5, color=BLUE, bold=True)
    doc.add_page_break()

    add_heading(doc, "보고서 요약", 1)
    add_callout(
        doc,
        "핵심 제안",
        "TCAD를 대체하는 것이 아니라 TCAD로 확보한 물리 기반 데이터를 Random Forest surrogate model로 학습하고, 넓은 공정 설계 공간을 빠르게 탐색한 뒤 최종 추천 레시피만 TCAD로 재검증한다.",
    )
    add_body(doc, "본 프로젝트는 planar bulk NMOS의 공정조건과 전기적 특성 사이의 비선형 관계를 학습하여, 사용자가 요구한 Vth, SS, Ion, Ioff에 가까운 공정 레시피를 제안하는 역설계 시스템을 구현하였다. 먼저 TCAD 공정조건 X에서 소자특성 Y를 예측하는 forward surrogate를 학습하고, UI에서는 학습 범위 안의 후보 공정조건 50만 개를 빠르게 평가하여 목표 Y에 가장 가까운 유효 레시피 X를 찾는다.")
    add_body(doc, "모델은 ‘좋은 소자/나쁜 소자’를 바로 나누지 않는다. TCAD failure 또는 extraction 누락처럼 정답을 신뢰하기 어려운 레시피를 Valid/Invalid classifier가 먼저 판별하고, 수치적으로 신뢰할 수 있는 레시피는 성능이 좋거나 나쁜 것과 관계없이 Vth·SS·Ion·Ioff 회귀기의 학습에 사용한다. 또한 채널 길이에 따른 물리 현상과 입력변수 차이를 반영하여 Short와 Long 모델을 분리하였다.")
    add_body(doc, "평가는 두 단계로 구분하였다. Forward test는 보지 않은 공정조건을 입력했을 때 스펙을 맞히는 X→Y 모델 자체의 성능을 확인한다. Inverse test는 목표 스펙을 UI에 입력해 얻은 Top 1 레시피를 실제 TCAD에 다시 넣고, 모델이 제시한 예상 스펙과 실제 스펙의 오차 및 최초 목표범위 충족 여부를 확인한다. Long은 두 평가에서 안정적인 결과를 보였고, Short는 Vth·SS·특히 Ioff에서 추가 데이터와 extraction 안정화가 필요한 것으로 나타났다.")
    add_body(doc, "Inverse TCAD 검증에서 Long 모델은 Vth MAE 10.2 mV, SS MAE 1.47 mV/dec, Ion log MAE 0.017 decade, Ioff log MAE 0.282 decade를 기록하였다. Short 모델은 각각 85.1 mV, 23.2 mV/dec, 0.125 decade, 1.785 decade로 나타나, Long 영역의 실용 가능성과 Short 영역의 보완 우선순위를 동시에 확인하였다.")

    add_heading(doc, "목차", 1)
    toc_items = [
        "서론 — 연구 배경, 문제 정의, 목표와 기여",
        "소자 구조 및 TCAD 데이터 구축 — 채널 구분, 공정변수, 문제 해결",
        "머신러닝 시스템 구조 — Random Forest, 역설계 알고리즘 및 사용자 UI",
        "데이터 처리 및 평가 설계 — Valid/Invalid, 전처리, 테스트셋",
        "성능 평가 및 결과 해석 — Forward/Inverse 성능과 물리적 해석",
        "결론 및 향후 과제",
        "부록 — 제출 전 확인사항 및 참고자료",
    ]
    for item in toc_items:
        p = doc.add_paragraph(style="List Number")
        # The visible section number is part of the report title, while the Word list remains real numbering.
        r = p.add_run(item)
        set_font(r, size=10.2, color=INK)
        p.paragraph_format.space_after = Pt(5)
    doc.add_page_break()

    add_heading(doc, "1. 서론", 1)
    add_heading(doc, "1.1 연구 배경", 2)
    add_body(doc, "반도체 소자의 전기적 특성은 gate length, implant dose와 energy, 열처리 조건 및 junction profile 등 여러 공정변수의 복합적인 영향을 받는다. Sentaurus TCAD는 이러한 관계를 물리 기반으로 해석할 수 있지만, 넓은 설계 공간의 모든 조합을 직접 계산하기에는 시간이 많이 소요된다. 특히 공정변수가 증가하면 가능한 조합 수가 급격히 커져 반복적인 설계 탐색의 부담이 커진다.")
    add_body(doc, "본 프로젝트는 TCAD 결과를 학습한 surrogate model을 이용하여 탐색 속도를 높이는 접근을 택하였다. 모델의 역할은 TCAD를 제거하는 것이 아니라, 많은 후보 중 유망한 공정조건을 빠르게 선별하여 고비용 TCAD 검증 횟수를 줄이는 것이다.")
    add_heading(doc, "1.2 문제 정의", 2)
    add_body(doc, "Forward 문제는 공정조건 X=[Lg, LDD Dose, LDD Energy, S/D Dose, S/D Energy, …]를 입력받아 Y=[Vth, SS, Ion, Ioff]를 예측하는 것이다. 그러나 실제 설계자는 원하는 소자 특성 Y를 먼저 제시하고 이를 만족하는 공정조건 X를 요구한다. 따라서 본 시스템은 학습된 X→Y 모델을 대규모 후보 탐색과 결합하여 Y→X 추천을 수행한다.")
    add_callout(doc, "연구 질문", "① 물리적으로 유효한 공정조건을 구분할 수 있는가? ② 요구 스펙에 가까운 레시피를 빠르게 추천할 수 있는가? ③ 추천값을 실제 TCAD에 재입력했을 때 예측 성능이 유지되는가?")
    add_heading(doc, "1.3 연구 목표와 기여", 2)
    add_table(doc, ["구분", "본 프로젝트의 기여"], [
        ["물리 기반 데이터", "Sentaurus Process/Device 결과를 학습 데이터로 사용"],
        ["채널별 모델", "Short/Long 물리 및 입력변수 차이를 반영한 독립 모델 bundle"],
        ["품질 제어", "Valid/Invalid 분류와 회귀 학습을 분리하여 실패 레시피 차단"],
        ["역설계", "Lg별 50만 개 후보 중 목표 제약과 오차를 고려해 Top 1 추천"],
        ["폐루프 검증", "추천 공정조건을 TCAD로 재계산하여 실제 오차 평가"],
    ], [2000, 7638])

    add_heading(doc, "2. 소자 구조 및 TCAD 데이터 구축", 1)
    add_heading(doc, "2.1 소자 구조와 기본 공정", 2)
    add_body(doc, "대상 소자는 simple planar bulk NMOS이다. 초기 p-type substrate 위에 gate oxide와 polysilicon gate를 형성하고, LDD implant, nitride spacer, main source/drain implant, annealing 및 aluminum contact 순서로 공정을 구성하였다. 전기적 특성은 Sentaurus Device의 ID–VG 해석으로 추출하였다.")
    add_table(doc, ["공정 단계", "주요 조건 또는 역할"], [
        ["Initial substrate", "Boron 1×10¹⁷ cm⁻³의 p-type bulk substrate"],
        ["Gate stack", "SiO₂ gate oxide 및 polysilicon gate"],
        ["LDD implant", "Arsenic dose/energy 변화"],
        ["Spacer", "Nitride spacer; geometry 일부는 Lg에 연동"],
        ["Main S/D implant", "Phosphorus dose/energy 변화"],
        ["Annealing", "Short 모델에서 thermal budget을 입력변수로 포함"],
        ["Contact", "Aluminum contact 형성 후 SDevice 해석"],
    ], [2700, 6938])
    add_heading(doc, "2.2 Short/Long 채널 분리 이유", 2)
    add_body(doc, "채널 길이가 짧아지면 drain 전계가 source 쪽 potential barrier에 영향을 주는 short-channel effect(SCE), Vth roll-off, DIBL 및 SS degradation이 두드러진다. 또한 Short 데이터에는 halo와 anneal 변수가 추가되어 Long과 입력 차원 자체가 다르다. 두 영역을 하나의 함수로 강제로 학습하면 Long의 안정적인 패턴이 Short의 급격한 비선형성과 섞일 수 있으므로, 모델과 후보 탐색 공간을 분리하였다.")
    add_table(doc, ["구분", "Lg 선택", "입력 공정변수", "모델 분리 근거"], [
        ["Short", "65, 90 nm", "Lg, LDD dose/energy, S/D dose/energy, anneal, halo dose/energy (8개)", "SCE·halo·thermal budget의 강한 비선형성"],
        ["Long", "180, 360, 720, 1000 nm", "Lg, LDD dose/energy, S/D dose/energy (5개)", "기본 공정변수 중심의 안정적인 특성"],
    ], [1300, 1800, 3838, 2700], font_size=8.3)
    add_heading(doc, "2.3 Short-channel 문제와 공정 보완", 2)
    add_body(doc, "초기 공정은 약 250 nm급 planar MOSFET을 기준으로 구성되어 있었기 때문에 gate length만 65/90 nm로 줄였을 때 물리적 scaling이 충분하지 않았다. 그 결과 일부 레시피에서 Vth extraction failure, 수천~수만 mV/dec의 SS, 매우 큰 leakage와 약한 gate response가 나타났다. 이때 모든 비정상 수치를 곧바로 ML outlier로 제거하지 않고, 먼저 소자 자체의 성능 열화인지 numerical·extraction failure인지를 구분하였다.")
    add_body(doc, "검토 결과 Short 영역에는 두 종류의 보완이 필요했다. 소자 측면에서는 SCE를 억제하기 위해 Boron halo dose/energy를 새 입력변수로 추가하고, gate oxide를 65 nm에서 2.0 nm, 90 nm에서 2.5 nm로 조정하였다. 데이터 측면에서는 ID–VG 곡선의 응답성과 Vth·SS 추출 성공 여부를 별도로 기록하여, 실제로 성능이 나쁜 소자와 계산 결과를 신뢰할 수 없는 레시피를 분리하였다.")
    add_table(doc, ["관찰된 문제", "원인 가설", "반영한 조치"], [
        ["Vth extraction 실패 또는 비정상 값", "SCE, sweep 범위, gate control 저하", "Curve validity와 extraction 결과를 분리 기록"],
        ["SS가 수천~수만 mV/dec", "실제 열화와 잘못된 extraction 지점이 혼재", "ID–VG 형상 및 추출 가능 여부를 함께 확인"],
        ["기존 250 nm급 recipe의 단순 축소", "oxide/junction/열처리 scaling 미반영", "65 nm 2.0 nm, 90 nm 2.5 nm oxide 적용"],
        ["Short에서 심한 SCE", "채널 중앙 barrier 제어 부족", "Boron halo dose/energy를 공정변수로 추가"],
    ], [2500, 3000, 4138], font_size=8.5)
    add_body(doc, "SVisual 결과에서도 정상 소자와 비정상 소자의 차이를 직접 확인하였다. 비정상 사례에서는 gate voltage가 증가해도 일반적인 subthreshold–turn-on 형태가 명확하지 않은 반면, 정상 사례에서는 log scale의 ID–VG 곡선에서 여러 decade에 걸친 전류 증가와 이후의 on-state 전이가 나타난다. NetActive 분포는 이러한 전기적 응답을 실제 공정으로 형성된 source/drain 및 channel 도핑 구조와 함께 검토하기 위해 사용하였다.")
    doc.add_picture(str(device_curve_figure), width=Cm(16.7))
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    add_figure_caption(doc, "그림 1. SVisual에서 확인한 비정상 gate-response 사례(상)와 정상 turn-on 사례(하)의 ID–VG 및 NetActive 분포. 두 ID–VG 곡선의 Y축은 log scale이다.")
    add_callout(doc, "중요한 데이터 원칙", "큰 SS나 낮은 Ion은 곧바로 numerical invalid가 아니다. TCAD가 정상 수렴하고 extraction이 물리적으로 의미 있는 구간에서 수행되었다면 ‘성능이 나쁜 소자’도 학습해야 한다. 반대로 solver failure 또는 잘못된 extraction으로 얻은 값은 invalid로 분리한다.", fill=PALE_GOLD, accent="7A5A00")
    add_heading(doc, "2.4 출력 특성의 bias 및 extraction 정의", 2)
    add_table(doc, ["출력", "추출 조건", "의미"], [
        ["Vth", "Vd = 0.05 V의 ID–VG", "낮은 drain bias에서 Vtgm 방식으로 추출한 threshold 특성"],
        ["SS", "Vd = 0.05 V의 ID–VG", "subthreshold 구간의 gate control을 나타내는 mV/dec 지표"],
        ["Ion", "Vd = 1.0 V, Vg = 2.5 V", "정의된 on-bias에서의 |Id|"],
        ["Ioff", "Vd = 1.0 V, Vg = 0 V", "정의된 off-bias에서의 |Id|"],
    ], [1800, 3100, 4738])
    add_callout(doc, "왜 추출 정의를 고정해야 하는가", "같은 이름의 Vth·SS·Ion·Ioff라도 bias와 추출 지점이 바뀌면 서로 다른 정답이 된다. 학습 데이터, UI 예측값, inverse TCAD 재검증에는 반드시 동일한 SVisual 정의를 적용해야 한다.")

    add_heading(doc, "3. 머신러닝 시스템 구조", 1)
    add_heading(doc, "3.1 머신러닝과 Random Forest 선정 이유", 2)
    add_heading(doc, "3.1.1 머신러닝이란 무엇인가", 3)
    add_body(doc, "머신러닝(Machine Learning)은 사람이 모든 계산식을 직접 정해 주는 대신, 입력과 정답이 함께 있는 여러 사례를 컴퓨터에 보여 주고 그 안의 반복되는 관계를 학습시키는 방법이다. 예를 들어 일반적인 프로그램이라면 연구자가 ‘Lg가 줄고 halo dose가 증가할 때 Vth가 얼마가 된다’는 식을 먼저 작성해야 한다. 그러나 실제 MOSFET에서는 implant dose·energy, anneal, halo와 Lg가 서로 영향을 주기 때문에 하나의 간단한 식으로 전체 관계를 표현하기 어렵다.")
    add_body(doc, "본 프로젝트에서 한 개의 학습 사례는 하나의 TCAD 공정 레시피이다. 입력 X에는 Lg, LDD 및 S/D implant 조건, anneal, halo 조건이 들어가고, 정답 Y에는 같은 레시피를 Sentaurus로 계산해 얻은 Vth, SS, Ion, Ioff가 들어간다. 모델은 수천 개의 X–Y 사례를 보면서 ‘어떤 공정영역에서 어떤 소자특성이 나타나는가’를 근사한다. 따라서 머신러닝 모델은 반도체 물리식을 없애는 장치가 아니라, TCAD로 계산한 관계를 빠르게 재현하는 surrogate model이다.")
    add_callout(doc, "이 프로젝트에서 머신러닝이 하는 일", "학습할 때는 공정조건 X → TCAD 소자특성 Y의 관계를 익힌다. 추천할 때는 목표 Y를 직접 X로 역산하지 않고, 많은 후보 X를 학습된 모델에 통과시킨 뒤 목표 Y와 가장 가까운 유효 후보를 고른다.")

    add_heading(doc, "3.1.2 왜 딥러닝이 아닌가", 3)
    add_body(doc, "딥러닝은 머신러닝과 별개의 개념이 아니라, 여러 hidden layer를 가진 인공신경망을 사용하는 머신러닝의 한 종류이다. 이미지·음성·문장 또는 매우 큰 원시 데이터처럼 특징을 사람이 미리 정하기 어려운 문제에서는 딥러닝이 강력하다. 반면 본 프로젝트의 데이터는 각 열의 물리적 의미가 이미 명확한 5개 또는 8개의 수치형 공정변수로 구성된 표 형식 데이터이며, 학습 가능한 레시피 수도 수천 개 규모이다.")
    add_table(doc, ["판단 기준", "본 프로젝트의 상황", "선택에 미친 영향"], [
        ["데이터 형태", "CSV/Excel의 수치형 tabular data", "트리 기반 모델에 적합"],
        ["입력 차원", "Long 5개, Short 8개 공정변수", "대규모 신경망 표현력이 필수적이지 않음"],
        ["데이터 비용", "한 레시피마다 TCAD 계산 필요", "많은 학습자료가 필요한 모델은 부담"],
        ["설명 가능성", "공정변수의 영향과 물리적 타당성 설명 필요", "분기와 SHAP을 활용할 수 있는 모델 선호"],
        ["개발 목적", "빠른 surrogate와 후보 탐색", "학습·튜닝·추론이 단순한 모델 선호"],
    ], [1900, 3900, 3838], font_size=8.5)
    add_body(doc, "따라서 딥러닝이 원리적으로 불가능해서 제외한 것은 아니다. 현재 데이터 조건에서는 신경망의 추가 복잡성이 반드시 더 높은 성능으로 이어진다고 보기 어렵고, 튜닝과 설명의 부담은 커진다. 향후 ID–VG 원시 곡선, 2차원 doping profile 또는 훨씬 많은 TCAD 사례를 직접 학습한다면 신경망을 비교 후보로 다시 검토할 수 있다.")

    add_heading(doc, "3.1.3 다른 알고리즘보다 Random Forest를 선택한 이유", 3)
    add_body(doc, "후보 알고리즘은 각각 장점이 있지만, 본 연구에서는 비선형성·표 형식 데이터·적은 전처리·설명 가능성·구현 안정성을 함께 고려해 Random Forest를 1차 surrogate로 선정하였다. 이는 Random Forest가 모든 경우에 가장 우수하다는 뜻이 아니라, 현재 데이터와 출품 시스템의 목적에 가장 균형 잡힌 선택이라는 의미이다.")
    add_table(doc, ["알고리즘", "장점", "현재 과제에서의 한계 또는 판단"], [
        ["선형회귀", "빠르고 계수 해석이 쉬움", "SCE와 공정변수 상호작용을 하나의 직선 관계로 표현하기 어려움"],
        ["k-NN", "가까운 기존 레시피를 직관적으로 참조", "변수 스케일에 민감하고 50만 후보를 반복 평가할 때 계산 부담이 큼"],
        ["SVR(RBF)", "비선형·중소규모 데이터에 강점", "scaling과 hyperparameter에 민감하며 결과 설명이 상대적으로 어려움"],
        ["신경망", "복잡하고 연속적인 함수 근사 가능", "현재의 소규모 tabular data에서는 과적합·튜닝·해석 부담이 큼"],
        ["Gradient Boosting", "tabular prediction에서 강력한 경쟁 모델", "순차 학습과 세밀한 튜닝이 필요하며 향후 비교 실험 대상으로 유지"],
        ["Random Forest", "비선형·상호작용 학습, scaling 부담이 낮고 안정적", "300개 트리 평균으로 단일 트리의 흔들림을 줄이고 병렬 추론 가능"],
    ], [1900, 3200, 4538], font_size=8.1)

    add_heading(doc, "3.1.4 Random Forest를 직관적으로 이해하기", 3)
    add_body(doc, "Random Forest는 여러 개의 의사결정 트리(Decision Tree)를 모아 하나의 예측을 만드는 모델이다. 이를 여러 사람이 각각 ‘스무고개’를 한 뒤 답을 종합하는 방식에 비유할 수 있다. 한 사람, 즉 트리 한 개는 입력된 공정 레시피에 예/아니오 질문을 차례로 던지면서 예상되는 소자특성의 범위를 좁혀 간다. 그리고 서로 다른 방식으로 학습한 여러 트리의 답을 모아 평균하거나 확률로 결합하므로, 한 트리의 판단에만 의존하는 것보다 결과가 안정적이다.")
    tree_intro = add_body(doc, "예를 들어 SS를 예측하는 트리의 첫 질문이 ‘Lg가 0.0775 µm 이상인가?’일 수 있고, 다음 질문은 ‘halo dose가 1.0×10¹³ cm⁻² 이상인가?’ 또는 ‘anneal time이 0.30 s 이하인가?’일 수 있다. 질문에 따라 학습 데이터가 점점 비슷한 그룹으로 나뉘고, 마지막 leaf에는 그 그룹의 평균 소자특성이 저장된다. 일반적인 스무고개와 달리 질문의 순서와 기준값을 사람이 미리 정하지는 않는다. 학습 알고리즘이 데이터에서 예측 오차를 가장 많이 줄이는 질문을 자동으로 선택한다.")
    tree_intro.paragraph_format.keep_with_next = True
    doc.add_picture(str(forest_figure), width=Cm(16.7))
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    add_figure_caption(doc, "그림 2. TCAD 공정변수로 구성한 Random Forest 개념도. 표시된 threshold와 SS 값은 이해를 위한 예시이며 실제 학습 트리를 그대로 출력한 값은 아니다.")
    add_body(doc, "한 개의 decision tree는 학습 데이터가 조금만 바뀌어도 질문 순서가 달라질 수 있다. Random Forest는 이를 보완하기 위해 서로 다른 bootstrap 표본으로 많은 트리를 만든다. 분류기는 각 트리가 계산한 valid 확률을 평균하고, 회귀기는 각 트리의 예측값을 평균한다. 현재 시스템은 모델 하나당 300개의 트리를 사용하므로, 한 공정조건의 최종 예측은 단일 규칙이 아니라 300개 서로 다른 판단의 결합 결과이다.")
    add_callout(doc, "그림을 읽는 방법", "이 그림은 SS 회귀기 내부의 ‘한 트리’와 ‘숲 전체’를 단순화한 것이다. 실제 시스템에서는 Vth, SS, Ion, Ioff마다 별도의 Random Forest가 있으며, Valid/Invalid 판단에도 별도의 Random Forest 분류기가 사용된다.", fill=PALE_GOLD, accent="7A5A00")
    add_heading(doc, "3.2 한 개의 트리가 학습되는 과정", 2)
    add_body(doc, "각 트리는 전체 학습행에서 복원추출한 bootstrap dataset을 사용한다. 현재 scikit-learn 기본 설정에서 분류기는 node마다 sqrt(전체 feature 수)만큼의 무작위 feature 후보를 검사하고 Gini impurity를 가장 많이 줄이는 split을 고른다. 회귀기는 node마다 모든 feature를 후보로 검사하되, bootstrap 표본이 트리마다 다르며 squared error를 가장 많이 줄이는 split을 고른다. max_depth를 별도로 제한하지 않았으므로 min_samples_split=2, min_samples_leaf=1 등 기본 종료조건에 도달할 때까지 성장한다.")
    add_table(doc, ["단계", "트리 내부에서 일어나는 일", "본 시스템의 기준"], [
        ["1. 표본 구성", "원래 학습자료에서 같은 개수만큼 복원추출", "트리마다 서로 다른 bootstrap 표본 사용"],
        ["2. 변수 후보", "현재 node에서 질문할 입력변수 후보를 정함", "분류기: √p개 무작위 후보, 회귀기: p개 전체"],
        ["3. 기준값 탐색", "후보 변수의 가능한 threshold를 비교", "분류: Gini 감소, 회귀: squared error 감소가 최대인 split"],
        ["4. 반복 분할", "좌·우 자료에 다시 질문을 만들어 내려감", "깊이 제한 없이 기본 종료조건까지 성장"],
        ["5. leaf 출력", "더 나누지 않는 마지막 그룹에서 답을 저장", "분류: Valid 비율, 회귀: 목표값 평균"],
        ["6. 숲의 결합", "각 트리의 답을 하나로 합침", "300개 확률 또는 예측값의 평균"],
    ], [1450, 4938, 3250], font_size=8.1)
    add_body(doc, "따라서 학습된다는 것은 단순히 값이 많이 몰린 쪽을 찾는 것이 아니라, 각 node에서 ‘어떤 변수를 어떤 기준값으로 질문하면 정답의 불순도 또는 예측오차가 가장 크게 감소하는가’를 반복해서 결정하는 과정이다. 분류기는 Valid/Invalid가 더 잘 갈라지도록, 회귀기는 같은 leaf 안의 Vth·SS·Ion·Ioff가 더 비슷해지도록 분기를 만든다.")
    add_callout(doc, "무작위인 부분과 학습되는 부분", "bootstrap 표본과 분류기 node에서 검토할 변수 집합은 무작위다. 그러나 그 후보 중 어느 변수와 기준값을 실제 질문으로 채택할지는 손실을 가장 많이 줄이는 결과로 결정된다. 변수 하나만 임의로 뽑는 것이 아니며, 한 트리에 모든 변수가 반드시 등장하는 것도 아니다.", fill=PALE_GOLD, accent="7A5A00")
    add_heading(doc, "3.3 Valid/Invalid 분류기–회귀기 2단계 구조", 2)
    add_heading(doc, "3.3.1 왜 분류기를 먼저 사용하는가", 3)
    add_body(doc, "TCAD 결과에는 서로 다른 두 종류의 ‘좋지 않은 데이터’가 섞일 수 있다. 첫째는 SS가 크거나 Ion이 낮은 것처럼 계산은 정상적으로 끝났지만 소자 성능이 나쁜 레시피이다. 이 자료는 ‘이 공정조건에서는 성능이 나빠진다’는 관계를 알려 주므로 버리면 안 된다. 둘째는 solver failure, Vth/SS extraction 실패 또는 필요한 출력 누락처럼 정답 자체를 신뢰하기 어려운 레시피이다. 이 값을 회귀기의 정답으로 넣으면 모델이 numerical failure를 실제 소자특성으로 오해할 수 있다.")
    add_body(doc, "따라서 본 시스템은 먼저 Random Forest classifier로 레시피가 수치 예측에 사용 가능한 Valid 영역인지 판단하고, 그 다음 Random Forest regressor로 소자특성의 크기를 예측한다. 분류기는 ‘성능이 좋은가’를 판단하는 모델이 아니라 ‘TCAD 결과와 extraction을 학습 정답으로 신뢰할 수 있는가’를 판단하는 품질 관문이다.")
    add_callout(doc, "Valid와 좋은 소자는 같은 뜻이 아니다", "큰 SS, 높은 Ioff 또는 낮은 Ion도 TCAD 곡선과 extraction이 정상이라면 Valid이다. 이러한 성능 열화 자료까지 회귀기가 학습해야 추천 과정에서 나쁜 공정영역을 실제로 피할 수 있다.", fill=PALE_GOLD, accent="7A5A00")

    add_heading(doc, "3.3.2 학습 단계와 추천 단계의 차이", 3)
    add_body(doc, "학습 단계에서는 classifier의 예측으로 자료를 거르는 것이 아니다. 이미 데이터에 기록된 실제 recipe-level Valid/Invalid label을 사용한다. 분류기는 Valid와 Invalid 전체 레시피를 보고 두 영역의 경계를 학습한다. 반면 Vth·SS·Ion·Ioff 회귀기는 실제 Valid이면서 네 출력이 모두 정상적으로 존재하는 레시피만 학습한다. 즉 Invalid 자료도 버려지는 것이 아니라 classifier를 학습시키는 중요한 negative example로 사용된다.")
    add_body(doc, "추천 단계에서는 사용자가 선택한 이산 Lg에서 후보 공정조건을 만든 뒤 classifier가 각 후보의 P(valid)를 계산한다. 현재 기준인 0.5보다 낮은 후보는 회귀기에 전달하지 않고 제외한다. P(valid)≥0.5인 후보만 네 개의 독립적인 regressor를 통과하여 Vth, SS, Ion, Ioff 예측값을 얻고, 이후 목표 스펙과의 오차를 계산한다.")
    doc.add_picture(str(two_stage_figure), width=Cm(16.7))
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    add_figure_caption(doc, "그림 3. 학습과 추천 단계에서의 Valid/Invalid classifier–regressor 연결 구조")

    add_heading(doc, "3.3.3 현재 Short/Long 모델 bundle", 3)
    add_table(doc, ["구성요소", "Short", "Long"], [
        ["Valid/Invalid 분류기", "RandomForestClassifier 300 trees", "RandomForestClassifier 300 trees"],
        ["Vth 회귀기", "RandomForestRegressor 300 trees", "RandomForestRegressor 300 trees"],
        ["SS 회귀기", "log₁₀(SS) 학습, 300 trees", "log₁₀(SS) 학습, 300 trees"],
        ["Ion 회귀기", "log₁₀(|Ion|) 학습, 300 trees", "log₁₀(|Ion|) 학습, 300 trees"],
        ["Ioff 회귀기", "log₁₀(|Ioff|) 학습, 300 trees", "log₁₀(|Ioff|) 학습, 300 trees"],
    ], [2400, 3619, 3619], font_size=8.5)
    doc.add_picture(str(flow_figure), width=Cm(16.7))
    add_figure_caption(doc, "그림 4. TCAD 데이터 생성부터 역설계 추천 및 TCAD 재검증까지의 전체 흐름")
    add_heading(doc, "3.4 역설계 추천 알고리즘", 2)
    add_heading(doc, "3.4.1 왜 목표 스펙에서 레시피를 직접 출력하는 Y→X 모델을 만들지 않았는가", 3)
    add_body(doc, "TCAD 데이터가 자연스럽게 제공하는 관계는 공정조건 X를 입력하면 소자특성 Y가 결정되는 X→Y 관계이다. 반대로 목표 Vth·SS·Ion·Ioff를 입력하여 공정 레시피를 바로 출력하는 Y→X 관계는 하나의 정답을 갖지 않는다. 서로 다른 LDD·S/D·halo·anneal 조합이 비슷한 소자특성을 만들 수 있기 때문에 forward 관점에서는 여러 X가 비슷한 Y로 모이는 many-to-one이고, 이를 역으로 풀면 하나의 Y에 여러 X가 대응하는 one-to-many 문제가 된다.")
    add_body(doc, "이 자료로 일반적인 회귀기를 Y→X 방향으로 학습하려면 동일한 목표에 대응하는 여러 레시피 중 하나를 임의의 정답으로 선택해야 한다. 또한 회귀기가 서로 다른 정답 레시피를 평균하면 학습 데이터에 실제로 존재하지 않거나 공정적으로 유효하지 않은 중간 조합을 출력할 수 있다. 특히 이산 Lg와 서로 다른 Short/Long 입력변수까지 동시에 다루면 이러한 모호성이 더 커진다. 따라서 직접 역회귀가 원리적으로 불가능한 것은 아니지만, 현재 데이터 규모와 설명 가능성을 고려하면 신뢰성 있는 단일 레시피를 보장하기 어렵다고 판단하였다.")
    add_callout(doc, "핵심 선택", "Y→X 함수를 직접 근사하는 대신, TCAD 데이터로 검증하기 쉬운 X→Y forward surrogate를 학습한 뒤 ‘어떤 X를 넣었을 때 목표 Y가 나오는가’를 빠르게 탐색한다.", fill=PALE_GOLD, accent="7A5A00")

    add_heading(doc, "3.4.2 X→Y 모델을 이용해 Y→X 추천을 수행하는 방법", 3)
    add_body(doc, "사용자가 목표 스펙 Y*와 허용범위를 입력하면 시스템은 공정조건 X를 미지수로 두고, 학습된 forward model의 예측값 Ŷ=f̂(X)가 Y*에 가장 가까워지는 X를 찾는다. 개념적으로는 X* = arg min S(f̂(X), Y*)를 푸는 과정이며, P(valid)≥0.5와 사용자가 지정한 스펙 범위를 제약조건으로 함께 적용한다.")
    add_body(doc, "먼저 사용자가 Short/Long과 학습 데이터에 존재하는 이산 Lg를 선택한다. 해당 Lg에서 나머지 공정변수의 학습 min–max 범위 안에 후보 레시피 500,000개를 생성한다. Dose 계열은 여러 자릿수 범위를 균형 있게 탐색하도록 log-uniform으로, energy와 anneal은 uniform으로 추출한다. 따라서 이는 범위를 무시한 무작위 찍기가 아니라, 학습 가능한 설계영역 안에서 수행하는 제한된 전역 탐색이다.")
    add_body(doc, "후보는 먼저 Valid/Invalid classifier를 통과한다. P(valid)<0.5인 후보는 제외하고, 남은 후보만 Vth·SS·Ion·Ioff 회귀기에 한꺼번에 입력하여 예상 특성 Ŷ를 계산한다. 사용자가 선택한 모든 min–max 범위를 만족하는 후보만 남긴 뒤, Vth·SS는 target 대비 상대오차 제곱, Ion·Ioff는 log₁₀ 공간의 오차 제곱을 합산한 score를 계산한다. score가 가장 작은 후보가 Top 1이며, UI에는 후속 선택을 위해 Top 3까지 표시한다.")
    doc.add_picture(str(inverse_design_figure), width=Cm(16.7))
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    add_figure_caption(doc, "그림 5. 직접 Y→X 회귀 대신 X→Y forward surrogate와 후보 탐색으로 역설계를 수행하는 방법")

    add_heading(doc, "3.4.3 50만 개 후보를 사용하는 이유와 의미", 3)
    add_body(doc, "Random Forest surrogate는 TCAD보다 매우 빠르게 대량 후보를 평가할 수 있으므로, 50만 개 후보도 UI에서 수 초 수준으로 비교할 수 있다. 후보 수를 늘리면 더 낮은 score의 레시피를 찾을 가능성은 높아지지만 계산시간과 메모리 사용량도 증가하고 개선폭은 점차 작아진다. 현재의 50만 개는 탐색 밀도와 응답시간을 절충한 설정이며, 난수 seed를 고정하여 같은 Lg에서는 동일한 후보 pool과 결과가 재현되도록 하였다.")
    add_body(doc, "이 방법이 수학적인 전역 최적해를 보장하는 것은 아니다. 그러나 여러 개의 가능한 레시피를 직접 비교할 수 있고, classifier와 공정범위 제약을 적용할 수 있으며, 각 후보의 예상 소자특성을 함께 제시할 수 있다는 장점이 있다. 향후에는 Latin hypercube·Sobol sampling, Bayesian optimization 또는 genetic algorithm을 결합하여 같은 계산량에서 탐색 효율을 높일 수 있다.")
    add_table(doc, ["단계", "처리 내용"], [
        ["1", "사용자가 Short/Long, 이산 Lg 및 목표 Vth·SS·Ion·Ioff 범위를 입력"],
        ["2", "해당 Lg에서 학습 데이터 min–max 내부의 공정후보 500,000개 생성"],
        ["3", "dose 계열은 log-uniform, energy/anneal은 uniform 방식으로 sampling"],
        ["4", "분류기 valid probability ≥ 0.5인 후보만 통과"],
        ["5", "4개 회귀기로 각 후보의 예상 스펙 계산"],
        ["6", "전류는 log error, Vth/SS는 상대오차 기반 score 계산"],
        ["7", "목표 범위를 만족하는 후보 중 score가 가장 작은 Top 1 추천"],
        ["8", "추천 레시피를 Sentaurus TCAD로 재검증"],
    ], [900, 8738])

    add_heading(doc, "3.5 사용자 UI 구성 및 사용 방법", 2)
    add_body(doc, "본 시스템은 역설계 알고리즘을 Flask 기반 웹 UI로 구현하였다. 사용자는 공정변수를 직접 추측하는 대신, 소자 영역과 gate length를 선택하고 원하는 Vth·SS·Ion·Ioff의 범위와 목표값을 입력한다. 서버는 동일한 조건에서 재현 가능한 50만 개 후보 pool을 평가하고, 목표와 가까운 공정 레시피를 Top 1–3으로 제시한다.")
    doc.add_picture(str(ui_usage_figure), width=Cm(16.7))
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    add_figure_caption(doc, "그림 6. 목표 스펙 입력부터 추천 공정 레시피와 예상 소자특성 확인까지의 실제 사용자 UI")

    add_heading(doc, "3.5.1 UI 사용 순서", 3)
    add_table(doc, ["순서", "사용자 조작", "시스템 내부 처리"], [
        ["1", "Short Channel 또는 Long Channel 선택", "서로 다른 입력변수와 학습자료를 사용한 해당 model bundle 선택"],
        ["2", "학습 데이터에 존재하는 이산 Lg 선택", "Short: 0.065/0.09 µm, Long: 0.18/0.36/0.72/1.0 µm 중 선택"],
        ["3", "사용할 지표를 체크하고 Min–Target–Max 입력", "min≤target≤max 확인 후 활성화된 지표만 제약조건과 score에 사용"],
        ["4", "‘최적 공정조건 탐색’ 버튼 클릭", "후보 생성 → valid probability 필터 → forward 예측 → 범위 필터 → score 정렬"],
        ["5", "Top 1–3 결과 확인", "추천 공정변수, 해당 레시피의 예상 Vth·SS·Ion·Ioff 및 valid probability 표시"],
        ["6", "Top 1 레시피를 Sentaurus TCAD로 재입력", "TCAD 실제값과 UI 예상값 및 최초 목표 스펙을 비교하여 최종 검증"],
    ], [850, 3750, 5038], font_size=8.2)

    add_heading(doc, "3.5.2 결과 화면을 해석하는 방법", 3)
    add_table(doc, ["표시 항목", "의미", "해석 시 주의점"], [
        ["Target spec", "사용자가 원하는 목표 Y*와 허용범위", "모델의 예측값이나 TCAD 정답이 아니라 탐색 기준"],
        ["Recommended process parameters", "Top 후보의 공정조건 X*", "Top 1은 현재 후보 pool에서 score가 가장 작은 조건"],
        ["Predicted device characteristics", "추천 X*를 forward model에 넣은 예상값 Ŷ=f̂(X*)", "Target과 같다고 가정하지 말고 두 값을 구분하여 확인"],
        ["Valid probability", "해당 후보가 numerical/extraction 관점에서 Valid일 분류 확률", "성능 만족 확률이나 예측 정확도·신뢰구간을 뜻하지 않음"],
        ["Top 1–3", "제약조건을 통과한 후보 중 score가 낮은 세 레시피", "공정 편의성과 TCAD 재검증 결과를 고려해 대안 선택 가능"],
    ], [2100, 3850, 3688], font_size=8.0)
    add_callout(doc, "UI에서 반드시 구분해야 하는 세 값", "① Target spec은 사용자가 요구한 목표 Y*, ② Predicted device characteristics는 추천 레시피에 대한 ML 예상값 Ŷ, ③ TCAD actual은 같은 레시피를 Sentaurus에 다시 넣어 얻는 최종 물리 기반 결과 Y_TCAD이다. 본 프로젝트의 inverse 평가는 ②와 ③의 차이를 정량화하고, 동시에 ③이 ①의 허용범위를 만족하는지도 확인한다.", fill=PALE_GOLD, accent="7A5A00")

    add_heading(doc, "3.6 SHAP 기반 모델 해석", 2)
    add_body(doc, "Random Forest가 높은 예측 성능을 보이더라도, 어떤 공정변수가 결과를 움직였는지 설명하지 못하면 semiconductor physics 관점의 검토가 어렵다. SHAP(SHapley Additive exPlanations)은 한 예측값을 기준 예측값과 각 입력변수의 기여도로 나누어 보여 주는 해석 방법이다. 예를 들어 특정 레시피에서 halo dose가 Vth 예측을 높이는 방향으로, Lg가 Ioff 예측을 낮추는 방향으로 기여했는지를 정량적으로 확인할 수 있다.")
    add_table(doc, ["해석 질문", "SHAP으로 확인하는 내용", "본 프로젝트에서의 활용"], [
        ["무엇이 중요한가?", "전체 데이터에서 |SHAP| 평균이 큰 변수", "Vth·SS·Ion·Ioff별 global feature importance 비교"],
        ["어느 방향으로 작용하는가?", "변수값의 증가가 예측을 높이거나 낮추는 방향", "halo, anneal, implant 조건의 물리적 경향 점검"],
        ["특정 추천은 왜 나왔는가?", "한 레시피 예측에 대한 변수별 기여", "추천 결과의 사후 설명과 이상 예측 점검"],
    ], [2300, 3600, 3738], font_size=8.4)
    add_body(doc, "SS·Ion·Ioff 회귀기는 log₁₀ 변환된 정답을 학습하므로 이 출력들의 SHAP 값도 모델 내부의 log 공간에 대한 기여도를 뜻한다. 따라서 SHAP 값의 부호와 크기를 실제 단위 변화로 해석할 때에는 10의 거듭제곱 관계를 함께 고려해야 한다.")
    add_callout(doc, "SHAP의 역할과 한계", "SHAP은 모델을 학습시키거나 Top 1을 선정하는 점수에 사용되지 않는다. 학습된 모델의 판단을 해석하는 도구이며, 변수 사이의 인과관계나 새로운 반도체 물리법칙을 증명하지는 않는다. TCAD 물리와 반대되는 SHAP 경향은 데이터 편향·상관관계·추출 오류를 점검하는 신호로 사용한다.")

    add_heading(doc, "4. 데이터 처리 및 평가 설계", 1)
    add_heading(doc, "4.1 학습 데이터 구성", 2)
    add_table(doc, ["구분", "원시 행", "공정 레시피", "Valid", "Invalid"], [
        ["Short", f"{len(short_sheet):,}", "5,072", "3,317", "1,755"],
        ["Long", f"{len(long_sheet):,}", "2,750", "2,332", "418"],
    ], [1800, 1800, 2000, 2000, 2038])
    add_body(doc, "한 공정 레시피는 Vd=0.05 V와 Vd=1.0 V의 두 행으로 구성된다. 모델 입력은 공정조건 기준으로 병합하고, Vth/SS는 low-Vd 행에서, Ion/Ioff는 high-Vd 행에서 선택하여 하나의 regression row를 만든다.")
    add_callout(doc, "행(row) 수와 레시피 수가 다른 이유", "Short 원시 데이터 10,144행은 서로 독립적인 10,144개 공정조건이 아니다. 동일 공정조건에 Vd=0.05 V와 1.0 V 결과가 각각 한 행씩 있으므로, 두 행을 합친 5,072개가 실제 recipe-level 학습 사례다.")
    add_heading(doc, "4.2 Valid/Invalid 처리", 2)
    add_body(doc, "Short 통합 데이터의 명시적 Valid flag를 authoritative label로 사용한다. 동일 레시피에서 두 drain bias가 모두 존재하고 두 행의 Valid가 모두 1이어야 recipe valid로 분류한다. Long과 같이 명시적 flag가 없는 기존 데이터는 두 bias에서 요구 출력이 모두 수치로 추출되었는지를 기준으로 label을 구성한다.")
    add_body(doc, "분류기는 valid/invalid 전체 레시피로 학습한다. 회귀기는 valid이면서 Vth, SS, Ion, Ioff가 모두 존재하고 Vth·SS·Ion이 양수인 레시피만 사용한다. 분류 단계와 회귀 단계를 분리함으로써 extraction failure를 숫자 0과 같은 가짜 성능값으로 학습시키지 않는다.")
    add_body(doc, "여기서 학습용 회귀자료를 고를 때는 classifier가 예측한 값이 아니라 CSV/Excel에 기록된 실제 Valid label을 사용한다. classifier의 예측 확률은 모델 학습이 끝난 뒤 UI가 새로운 후보를 평가할 때만 품질 필터로 사용한다. 즉 학습 정답과 모델 예측을 서로 섞지 않았다.")
    add_body(doc, "실제 전처리에서는 Sentaurus Workbench에서 SVisual node가 빨간색 fail로 표시되거나 Vth·SS 등 요구 출력이 누락된 경우를 확인한 뒤, CSV/Excel 통합표에서 해당 공정 레시피를 Valid=0으로 표시하였다. 정상적으로 계산·추출된 레시피는 Valid=1로 표시하여 classifier가 학습할 명시적 정답 label을 구성하였다.")
    doc.add_picture(str(validity_figure), width=Cm(16.7))
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    add_figure_caption(doc, "그림 7. Sentaurus TCAD의 fail 및 출력 누락을 CSV/Excel의 Valid/Invalid label로 변환한 실제 전처리 과정")
    add_callout(doc, "전처리 판정의 핵심", "빨간색 fail, extraction 실패 및 필수 출력 누락은 Invalid로 처리한다. 반면 계산과 extraction이 정상이라면 SS가 크거나 Ion이 낮은 성능 열화 레시피도 Valid로 유지하여 회귀기가 나쁜 공정영역까지 학습하도록 한다.", fill=PALE_GOLD, accent="7A5A00")
    add_callout(doc, "CurveValid의 역할", "CurveValid는 Vg 증가에 대한 ID–VG 곡선의 기본 응답성과 extraction 신뢰성을 확인하기 위한 보조 신호다. CurveValid=1이 곧 좋은 소자라는 뜻은 아니며, 최종 학습에서는 통합 데이터의 recipe-level Valid label이 분류 정답으로 사용된다.")
    add_heading(doc, "4.3 수치 변환", 2)
    add_body(doc, "Ion과 Ioff는 여러 order에 걸쳐 변화하므로 log₁₀(|I|) 공간에서 학습한다. SS 역시 극단값의 영향을 줄이기 위해 log₁₀(SS) 공간에서 학습하고, UI 출력 시 10의 거듭제곱으로 원 단위에 복원한다. Vth는 원 단위(V) 그대로 학습한다.")
    add_callout(doc, "왜 전류를 log로 학습하는가", "예를 들어 10⁻¹² A와 10⁻⁶ A는 숫자 차이가 작아 보이지만 물리적으로는 100만 배 차이다. 원 단위의 squared error를 그대로 사용하면 큰 전류가 손실을 지배할 수 있다. log 변환은 각 order의 상대적 차이를 비교 가능하게 만든다.", fill=PALE_GOLD, accent="7A5A00")
    add_heading(doc, "4.4 Forward 테스트: X→Y 예측 정확도 평가", 2)
    add_body(doc, "Forward test는 공정조건 X를 모델에 입력했을 때 Vth·SS·Ion·Ioff라는 소자특성 Y를 얼마나 정확히 예측하는지를 평가한다. 전체 레시피의 80%를 training set, 20%를 한 번도 학습에 사용하지 않은 hold-out test set으로 분리하고(random_state=42), 20%의 공정조건을 학습된 모델에 입력하여 얻은 예측값 Ŷ와 해당 레시피의 실제 TCAD값 Y를 비교하였다. 여기서 계산한 R²·MAE·RMSE는 순수한 X→Y surrogate model의 예측 성능을 의미한다.")
    add_body(doc, "분류 데이터는 valid/invalid 비율이 유지되도록 stratified split을 적용하였다. 회귀기는 네 출력이 완전한 valid recipe에 공통된 index split을 사용하여 동일한 테스트 레시피에서 Vth, SS, Ion, Ioff를 평가하였다. 성능 평가가 끝난 뒤 실제 UI에 사용되는 최종 모델은 이용 가능한 전체 데이터로 다시 학습하였다.")
    add_callout(doc, "80:20의 정확한 의미", "80:20은 정확도가 80%라는 뜻이 아니다. 전체 데이터 중 80%로 모델의 질문과 기준값을 학습하고, 정답을 보여 주지 않은 나머지 20%에서 예측오차를 측정했다는 데이터 분할 비율이다.")
    add_callout(doc, "Forward test만으로는 충분하지 않다", "Forward test는 보지 않은 공정조건의 스펙을 잘 예측하는지 보여 주지만, 사용자가 목표 스펙을 입력했을 때 선택되는 추천 레시피가 실제 TCAD에서도 그 목표를 달성하는지는 직접 검증하지 않는다. 따라서 본 프로젝트의 최종 목적에는 별도의 inverse test가 필요하다.", fill=PALE_GOLD, accent="7A5A00")

    add_heading(doc, "4.5 Inverse recommendation 테스트: Y→X→TCAD 전체 시스템 검증", 2)
    add_body(doc, "본 프로젝트가 궁극적으로 원하는 기능은 사용자가 UI에 목표 소자특성 Y*를 입력하면 이를 만족할 공정조건 X*를 추천하는 것이다. 따라서 inverse test는 UI의 입력부터 최종 TCAD 확인까지 전체 추천 경로를 평가한다. 먼저 실사용 가능성이 높은 영역에서 목표 Vth·SS·Ion·Ioff와 허용범위를 구성하고, 실제 UI와 동일하게 Lg별 500,000개 후보를 탐색하여 Top 1 공정 레시피 X*를 얻었다.")
    add_body(doc, "다음으로 추천된 X*를 Sentaurus TCAD에 직접 입력하여 실제 소자특성 Y_TCAD(X*)를 다시 추출하였다. 이 결과로 두 가지를 확인한다. 첫째, UI가 X*에 대해 함께 제시한 predicted device characteristics Ŷ(X*)와 실제 TCAD값이 얼마나 일치하는지를 MAE·RMSE·R²로 평가한다. 둘째, 실제 TCAD값이 사용자가 처음 요청한 min–max 목표범위 안에 들어오는지를 확인한다. 전자는 추천 지점에서의 surrogate 신뢰도이고, 후자는 최종적인 목표 스펙 달성 여부이다.")
    add_body(doc, "테스트 목표는 학습 데이터에서 실사용 가능성이 높은 출력 영역을 고른 뒤 약간 perturb하여 표에 그대로 존재하지 않는 값으로 구성하였다. 고정 seed 20260823과 Lg별 동일 후보 pool을 사용하여 UI와 batch 추천 결과가 재현되도록 하였다.")
    doc.add_picture(str(evaluation_design_figure), width=Cm(16.7))
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    add_figure_caption(doc, "그림 8. Forward model 평가와 실제 UI 목적을 검증하는 inverse recommendation 평가의 차이")
    add_callout(doc, "Inverse test의 의미", "목표 Y*를 입력했을 때 추천된 X*를 실제 TCAD에 다시 넣어 Y_TCAD(X*)를 얻고, 모델 예측과의 오차 및 최초 목표범위 충족 여부를 동시에 확인하는 end-to-end 검증이다.", fill=LIGHT_BLUE, accent=BLUE)
    add_table(doc, ["구분", "Lg별 케이스", "총 목표", "고유 Top 1 레시피", "TCAD 매칭"], [
        ["Short", "65/90 nm × 50", "100", "99", "99/99"],
        ["Long", "180/360/720/1000 nm × 50", "200", "175", "175/175"],
    ], [1600, 3000, 1500, 2000, 1538])

    add_heading(doc, "5. 성능 평가 및 결과 해석", 1)
    add_heading(doc, "5.1 평가 지표", 2)
    add_body(doc, "출력마다 수치 범위와 물리적 의미가 다르므로 하나의 ‘정확도 %’로 합치지 않고, 실제 단위의 오차와 설명력, 전류의 배수오차를 함께 제시하였다. Vth와 SS는 설계자가 직접 사용하는 V와 mV/dec 단위의 MAE·RMSE가 직관적이며, 여러 order에 걸쳐 변하는 Ion·Ioff는 log-domain 지표가 더 적합하다.")
    add_table(doc, ["지표", "정의와 해석", "적용 출력"], [
        ["R²", "실제값 변동 중 모델이 설명한 비율. 1에 가까울수록 좋으나 ‘정확도 %’와 동일하지 않음", "Vth, SS 및 log current"],
        ["MAE", "|예측−실제|의 평균. 실제 단위로 평균적인 오차 크기를 직관적으로 표현", "Vth, SS"],
        ["RMSE", "오차 제곱 평균의 제곱근. 큰 오차에 더 큰 패널티를 부여", "Vth, SS"],
        ["log MAE/RMSE", "log₁₀ 공간의 오차. 1 decade는 10배, 0.301 decade는 약 2배", "Ion, Ioff"],
    ], [1700, 5700, 2238], font_size=8.7)
    add_callout(doc, "지표를 읽는 예", "R²=0.99를 ‘99% 정확도’라고 부르지는 않는다. Vth MAE=0.0062 V는 테스트 레시피에서 평균 절대오차가 6.2 mV였다는 뜻이다. 전류 log MAE=0.301 decade는 예측값과 실제값이 기하평균 관점에서 약 10⁰·³⁰¹≈2배 차이임을 뜻한다.")
    add_heading(doc, "5.2 Forward prediction 성능", 2)
    add_table(doc, ["모델", "Vth", "SS", "Ion", "Ioff"], [
        ["Short", "R² 0.960\nMAE 28.0 mV", "R² 0.831\nlog RMSE 0.207 dec", "R² 0.987\nlog RMSE 0.609 dec", "R² 0.996\nlog RMSE 0.605 dec"],
        ["Long", "R² 0.986\nMAE 6.20 mV", "R² 0.994\nMAE 2.66 mV/dec", "R² 0.994\nlog RMSE 0.010 dec", "R² 0.794\nlog RMSE 0.267 dec"],
    ], [1200, 2100, 2200, 2069, 2069], font_size=8.2)
    add_body(doc, "Valid/Invalid 분류 정확도는 Short 97.93%(balanced accuracy 98.35%, ROC-AUC 99.93%), Long 99.45%(balanced accuracy 99.19%, ROC-AUC 99.92%)였다. Long은 전반적으로 안정적인 forward 성능을 보였다. Short는 Vth 설명력은 높지만, SS와 current가 넓은 동적 범위와 경계 레시피의 영향을 받아 log 오차가 상대적으로 크다.")
    add_body(doc, "높은 R²와 상대적으로 큰 log RMSE가 동시에 나타날 수 있다는 점도 중요하다. R²는 테스트셋 전체 변화 폭을 모델이 얼마나 따라가는지 평가하고, RMSE는 개별 레시피의 큰 오차에 민감하다. 따라서 Short Ion·Ioff의 높은 R²는 전체적인 증가·감소 경향을 잘 학습했다는 뜻이지만, 일부 경계 레시피에서 수 배 이상의 전류오차가 없다는 뜻은 아니다. 또한 Valid/Invalid 비율이 완전히 같지 않으므로 단순 accuracy뿐 아니라 두 class의 recall을 균형 있게 반영한 balanced accuracy와 ROC-AUC를 함께 제시하였다.")
    add_heading(doc, "5.3 Inverse recommendation 성능", 2)
    sm = short_metrics["summary"]["prediction_metrics"]
    lm = long_metrics["summary"]["prediction_metrics"]
    add_table(doc, ["모델", "Vth", "SS", "Ion", "Ioff"], [
        ["Short", f"R² {sm['Vth']['r2']:.3f}\nMAE {sm['Vth']['mae']*1000:.1f} mV\nRMSE {sm['Vth']['rmse']*1000:.1f} mV", f"R² {sm['SS']['r2']:.3f}\nMAE {sm['SS']['mae']:.1f}\nRMSE {sm['SS']['rmse']:.1f} mV/dec", f"log R² {sm['Ion']['log_r2']:.3f}\nlog MAE {sm['Ion']['log_mae']:.3f}\nlog RMSE {sm['Ion']['log_rmse']:.3f}", f"log R² {sm['Ioff']['log_r2']:.3f}\nlog MAE {sm['Ioff']['log_mae']:.3f}\nlog RMSE {sm['Ioff']['log_rmse']:.3f}"],
        ["Long", f"R² {lm['Vth']['r2']:.3f}\nMAE {lm['Vth']['mae']*1000:.1f} mV\nRMSE {lm['Vth']['rmse']*1000:.1f} mV", f"R² {lm['SS']['r2']:.3f}\nMAE {lm['SS']['mae']:.2f}\nRMSE {lm['SS']['rmse']:.2f} mV/dec", f"log R² {lm['Ion']['log_r2']:.3f}\nlog MAE {lm['Ion']['log_mae']:.3f}\nlog RMSE {lm['Ion']['log_rmse']:.3f}", f"log R² {lm['Ioff']['log_r2']:.3f}\nlog MAE {lm['Ioff']['log_mae']:.3f}\nlog RMSE {lm['Ioff']['log_rmse']:.3f}"],
    ], [1200, 2100, 2200, 2069, 2069], font_size=8.0)
    add_body(doc, "Inverse 성능은 모델이 추천한 공정조건에 대해 표시한 predicted device characteristics와 동일 공정을 TCAD에 투입해 얻은 실제 특성을 직접 비교한 결과다. 따라서 forward 20% hold-out보다 실제 시스템 사용 성능을 더 직접적으로 반영한다.")
    add_body(doc, "배수오차로 환산하면 Short의 Ion log MAE 0.125 decade는 약 10⁰·¹²⁵≈1.33배, Ioff 1.785 decade는 약 61배의 평균적인 차이를 뜻한다. Long은 Ion 0.017 decade가 약 1.04배, Ioff 0.282 decade가 약 1.91배이다. 이에 따라 Long 추천은 모델 예측과 TCAD 재계산이 매우 가까웠지만, Short Ioff는 추천값을 최종 확정하기 전에 TCAD 확인이 반드시 필요한 출력으로 판단하였다.")
    add_body(doc, "Short SS의 inverse R²가 낮은 것은 MAE 23.2 mV/dec와 모순되지 않는다. 해당 practical test가 비교적 좁은 SS 영역을 대상으로 구성되어 실제값의 분산이 작았기 때문에, 수십 mV/dec의 오차만으로도 분산 대비 설명력인 R²가 크게 낮아질 수 있다. 좁은 운용영역에서는 R² 하나보다 실제 단위의 MAE·RMSE를 함께 보는 것이 적절하다.")
    add_table(doc, ["모델", "Vth 목표범위", "SS 목표범위", "Ion 목표범위", "Ioff 목표범위", "4개 동시"], [
        ["Short", "53%", "84%", "90%", "17%", "9%"],
        ["Long", "98.5%", "100%", "100%", "76%", "76%"],
    ], [1200, 1688, 1688, 1688, 1688, 1686], font_size=8.2)
    add_body(doc, "목표범위 만족률은 테스트셋의 난이도와 허용범위에 종속되므로 일반적인 ‘정확도’로 표현해서는 안 된다. 위 표는 현재 practical target suite에서 추천 레시피가 각 요청 범위에 실제로 들어왔는지를 보여주는 보조지표다.")
    add_callout(doc, "Inverse 결과에서 답하는 두 질문", "① predicted device characteristics와 TCAD actual이 가까운가: 추천 지점의 surrogate 신뢰도. ② TCAD actual이 최초 min–max 범위에 들어오는가: 사용자의 목표 달성 여부. 첫 번째는 MAE·RMSE·R²로, 두 번째는 범위 만족률로 평가하며 서로 대체할 수 없다.", fill=PALE_GOLD, accent="7A5A00")
    add_heading(doc, "5.4 Short 성능 저하와 SS–Ioff 연관성", 2)
    add_body(doc, "Short inverse test에서 |SS 예측오차|와 |Ioff log 예측오차|의 Spearman 상관계수는 ρ=0.586(p<0.001)이었다. 이는 SS 예측이 크게 빗나간 레시피에서 Ioff도 함께 여러 배 또는 여러 order 어긋나는 경향이 있음을 뜻한다. 인과관계를 직접 증명하지는 않지만, 두 출력이 subthreshold ID–VG 곡선의 동일한 기울기와 위치 변화에 민감하다는 물리적 해석과 일치한다.")
    add_callout(doc, "왜 작은 SS 오차가 큰 Ioff 오차로 이어질 수 있는가", "Subthreshold에서 SS를 V/dec 단위로 놓으면 전류는 대략 I_D ∝ 10^((V_G−V_th)/SS)로 변화한다. 보고서의 SS가 mV/dec이면 먼저 1000으로 나누어 V/dec로 환산한다. 따라서 SS 또는 Vth의 작은 예측오차도 지수항을 통해 Ioff의 배수오차로 증폭될 수 있다. Short의 Ioff 개선에는 Ioff 모델만 조정하기보다 ID–VG extraction 일관성, SS·Vth 공동오차 및 경계 데이터 밀도를 함께 개선해야 한다.")
    add_heading(doc, "5.5 결과의 한계", 2)
    add_table(doc, ["한계", "보고서에서의 해석"], [
        ["Random split의 낙관성", "같은 Lg·유사 DOE 조합이 train/test에 함께 존재할 수 있음"],
        ["Lg extrapolation 취약", "Leave-one-Lg-out 성능이 낮아 UI에서 학습된 이산 Lg만 추천"],
        ["Short 경계 급변", "halo·anneal·oxide·junction 상호작용 근처에서 데이터 밀도가 부족"],
        ["RF 불확실성 미표시", "Top 1의 낮은 surrogate score가 실제 TCAD 오차를 보장하지 않음"],
        ["Extraction 정의 의존", "Vth/SS/Ion/Ioff 정의가 바뀌면 데이터 전체의 일관성을 다시 확인해야 함"],
    ], [2600, 7038])

    add_heading(doc, "6. 결론 및 향후 과제", 1)
    add_heading(doc, "6.1 결론", 2)
    add_body(doc, "본 프로젝트는 planar NMOS의 공정조건 X와 전기적 특성 Y의 관계를 학습하는 Short/Long Random Forest surrogate를 구축하고, 이를 목표 스펙 기반 역설계 시스템으로 확장하였다. 먼저 Valid/Invalid classifier가 numerical·extraction failure 가능성이 높은 후보를 차단하고, 네 개의 regressor가 유효 후보의 Vth·SS·Ion·Ioff를 예측한다. 이후 학습 범위 안의 후보 50만 개를 비교하여 목표에 가장 가까운 Top 1 레시피를 제시하고, 최종 후보는 다시 TCAD로 확인한다.")
    add_body(doc, "이 구조는 ‘나쁜 성능’과 ‘신뢰할 수 없는 계산’을 구분하여 실패 데이터도 분류 학습에 활용하고, 성능이 나쁘지만 정상인 레시피는 회귀 학습에 유지한다. 또한 Short/Long의 물리 현상과 입력변수 차이를 모델에 반영하고, 전류의 넓은 동적 범위를 log-domain으로 처리하였다.")
    add_body(doc, "Forward 평가는 surrogate 자체가 보지 않은 X에서 Y를 재현하는지 확인하는 필수 시험이며, Inverse 평가는 사용자가 요청한 Y*에서 추천 X*를 얻어 실제 TCAD의 Y가 나오는지 확인하는 최종 시스템 시험이다. Long 영역은 두 평가 모두 높은 재현성을 보였다. Short 영역은 halo와 oxide scaling 이후 SS 성능이 개선되었으나 Vth·SS·Ioff의 결합오차와 경계 민감성이 남아, 추가 DOE와 extraction 표준화가 우선 과제로 확인되었다.")
    add_heading(doc, "6.2 향후 개선", 2)
    add_table(doc, ["우선순위", "개선 항목", "기대 효과"], [
        ["1", "65/90 nm 경계 및 급변 영역 중심의 추가 DOE/active learning", "적은 TCAD 추가 계산으로 Short 오차 집중 감소"],
        ["2", "SS·Vth extraction 정의와 ID–VG 품질검사 표준화", "Ioff까지 연결되는 label noise 감소"],
        ["3", "Leff/Lg, Xj, lateral diffusion 등 physics-informed feature 추가", "SCE 원인 설명력과 extrapolation 개선"],
        ["4", "트리 간 분산 또는 conformal prediction 기반 불확실성 표시", "Top 1 과신 방지 및 재검증 우선순위 제공"],
        ["5", "Bayesian optimization/NSGA-II와 surrogate 결합", "50만 개 random search보다 효율적인 다목적 탐색"],
    ], [1200, 5300, 3138], font_size=8.6)
    add_callout(doc, "최종 메시지", "본 시스템의 가치는 TCAD를 ML로 치환하는 데 있지 않다. TCAD의 물리적 신뢰성과 ML의 탐색 속도를 결합하여, 설계자가 더 넓은 공정 공간을 빠르게 검토하고 최종 후보에 계산 자원을 집중할 수 있도록 하는 데 있다.")

    add_heading(doc, "부록 A. 제출 전 확인사항", 1)
    add_table(doc, ["항목", "확인 또는 입력할 내용"], [
        ["표지 정보", "출품 분야, 팀명, 작성자, 소속, 제출일"],
        ["최종 SVisual 정의", "Vth/SS extraction 식 최종 확인, Ion=|Id|(Vd=1 V, Vg=2.5 V), Ioff=|Id|(Vd=1 V, Vg=0 V)"],
        ["공정조건", "oxide 65 nm=2.0 nm, 90 nm=2.5 nm 및 최종 anneal 범위 재확인"],
        ["단위", "Lg(µm/nm), dose(cm⁻²), energy(keV), current(A 또는 A/µm) 통일"],
        ["그림", "정상/비정상 사례의 정의와 최종 그림 캡션 확인"],
        ["참고문헌", "사용한 Sentaurus 버전의 공식 매뉴얼과 반도체 소자 교재 명시"],
    ], [2500, 7138])
    add_heading(doc, "부록 B. 참고자료 초안", 1)
    refs = [
        "L. Breiman, ‘Random Forests,’ Machine Learning, vol. 45, pp. 5–32, 2001.",
        "scikit-learn Developers, ‘Ensembles: Gradient boosting, random forests, bagging, voting, stacking,’ scikit-learn User Guide, https://scikit-learn.org/stable/modules/ensemble.html (accessed 2026-08-25).",
        "Google for Developers, ‘Machine Learning Crash Course: Data characteristics and Neural networks,’ https://developers.google.com/machine-learning/crash-course/ (accessed 2026-08-25).",
        "Synopsys, Sentaurus Process / Sentaurus Device User Guide, 사용 버전 확인 후 기입.",
        "Y. Taur and T. H. Ning, Fundamentals of Modern VLSI Devices.",
        "S. M. Sze and K. K. Ng, Physics of Semiconductor Devices.",
        "프로젝트 내부자료: tcad_training_master.xlsx, inverse_test_results.xlsx, evaluation_short/long.json.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        r = p.add_run(ref)
        set_font(r, size=8.8, color=INK)
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.55)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(2)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    embed_font(
        OUT,
        ROOT / "reports" / "assets" / "Pretendard-Regular.ttf",
        FONT,
        bold_path=ROOT / "reports" / "assets" / "Pretendard-Bold.ttf",
        alt_name="Pretendard",
    )
    print(OUT)


if __name__ == "__main__":
    build_report()
